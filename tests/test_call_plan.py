from openpyxl import load_workbook

from analytics import call_plan
from analytics.call_plan import build_rows, read_carry


def plan(**kw):
    p = {"reviewed": [], "tiered": [], "first": [], "candidates": [], "stages": {}, "owner_order": ["Asha K", "Ravi S"]}
    p.update(kw)
    return p


def cand(lead, owner="Asha K", missed=0, stage="Call Back Later"):
    return {"lead_id": lead, "owner": owner, "name": f"Lead {lead}", "phone": "9", "stage": stage,
            "inbound_missed_unreturned": missed, "last_conversation_ist": "2026-10-06 10:00"}


def tiered(lead, tier, owner="Asha K", prob=10):
    return {"lead_id": lead, "owner": owner, "name": f"Lead {lead}", "tier": tier, "prob_3d": prob, "why": "w",
            "opening_line": "hi", "next_action": "n"}


def test_missed_call_from_lead_becomes_tier_m_and_goes_first():
    p = plan(candidates=[cand("L1", missed=2), cand("L2"), cand("L3", missed=1)],
             tiered=[tiered("L2", "A", prob=40), tiered("L3", "B")])
    rows = build_rows(p)["Asha K"]
    assert [r["lead_id"] for r in rows] == ["L3", "L1", "L2"]  # within M, the higher chance first
    assert rows[0]["tier"] == rows[1]["tier"] == "M" and "Called us 2 time(s)" in rows[1]["why"]
    assert rows[0]["opening_line"] == "hi"  # the reviewed line is kept
    assert rows[1]["opening_line"].startswith("Hi, this is Asha from LawSikho")  # the caller's name, not <name>


def test_link_sent_not_paid_carried_from_yesterdays_workbook(tmp_path):
    p = plan(candidates=[cand("L1"), cand("L2"), cand("L3")],
             tiered=[tiered("L1", "B"), tiered("L2", "C"), tiered("L3", "B")])
    by_owner = build_rows(p)
    wb_path = str(tmp_path / "wed.xlsx")
    wb = call_plan.Workbook()
    call_plan.guide_sheet(wb)
    _, first, _ = call_plan.caller_sheet(wb, "Asha K", by_owner["Asha K"])
    ws = wb["Asha K"]
    col = call_plan.COL
    ids = {ws.cell(row=first + i, column=col["Lead ID"]).value: first + i for i in range(3)}
    ws.cell(row=ids["L1"], column=col["Payment status"], value="Link sent")
    ws.cell(row=ids["L2"], column=col["Payment status"], value="Token paid")
    ws.cell(row=ids["L2"], column=col["Attempt 1 outcome"], value="Enrolled / paid")  # paid after all: not carried
    ws.cell(row=ids["L3"], column=col["Payment status"], value="Paid")
    wb.save(wb_path)

    carry = read_carry([wb_path])
    assert [c["lead_id"] for c in carry] == ["L1"] and carry[0]["owner"] == "Asha K"
    rows = build_rows(plan(candidates=[cand("L1")], carry=carry))["Asha K"]
    assert rows[0]["tier"] == "P" and "Link sent" in rows[0]["why"]


def test_summary_sheet_counts_new_tiers(tmp_path):
    rows = build_rows(plan(candidates=[cand("L1", missed=1)], tiered=[tiered("L2", "A")]))
    wb = call_plan.Workbook()
    call_plan.guide_sheet(wb)
    _, f, l = call_plan.caller_sheet(wb, "Asha K", rows["Asha K"])
    call_plan.summary_sheet(wb, ["Asha K"], {"Asha K": (f, l)})
    path = str(tmp_path / "p.xlsx")
    wb.save(path)
    ws = load_workbook(path)["Team summary"]
    heads = [c.value for c in ws[4]]
    assert heads[1:3] == ["Return calls (M)", "Unpaid links (P)"]
    assert ws.cell(row=5, column=2).value.startswith('=COUNTIF(') and '"M ·*"' in ws.cell(row=5, column=2).value


def test_tier_b_cap_moves_lowest_ranked_extras_to_least_loaded_caller():
    p = plan(candidates=[cand(f"B{i}") for i in range(5)] + [cand("A1")],
             tiered=[tiered(f"B{i}", "B", prob=20 - i) for i in range(5)] + [tiered("A1", "A")], tier_b_cap=3)
    rows = build_rows(p)
    asha = [r["lead_id"] for r in rows["Asha K"]]
    ravi = [r["lead_id"] for r in rows["Ravi S"]]
    assert asha == ["A1", "B0", "B1", "B2"] and sorted(ravi) == ["B3", "B4"]
    assert all(r["owner"] == "Ravi S" and r["why"].startswith("Moved from Asha K") for r in rows["Ravi S"])


def test_sheet_cap_keeps_hot_leads_and_moves_the_rest_to_the_backlog(tmp_path):
    p = plan(candidates=[cand(x) for x in ("A1", "B1", "C1", "C2", "C3")],
             tiered=[tiered("A1", "A"), tiered("B1", "B"), tiered("C1", "C"), {**tiered("C2", "C"), "rank": 9},
                     tiered("C3", "C")], sheet_cap=3)
    out = str(tmp_path / "plan.xlsx")
    rows = call_plan.write_workbook(p, out)["Asha K"]
    assert [r["lead_id"] for r in rows] == ["A1", "B1", "C2"]  # the best-ranked C keeps the last place
    wb = load_workbook(out)
    assert "Backlog" in wb.sheetnames
    assert sorted(wb["Backlog"].cell(row=r, column=10).value for r in (5, 6)) == ["C1", "C3"]
    p = plan(candidates=[cand(f"B{i}") for i in range(4)], tiered=[tiered(f"B{i}", "B") for i in range(4)], sheet_cap=2)
    assert len(build_rows(p)["Asha K"]) == 4 and call_plan.split_backlog(build_rows(p), 2) == {}  # B never cut
    p = plan(candidates=[cand("C1"), cand("C2")], tiered=[tiered("C1", "C", prob=9), {**tiered("C2", "C"), "month_end_due": True}])
    rows = build_rows(p)
    assert [r["lead_id"] for r in call_plan.split_backlog(rows, 1)["Asha K"]] == ["C1"]  # the due promise stays


def test_counsellors_leads_ready_to_close_go_to_closers():
    p = plan(owner_order=["Asha K", "Ravi S"],
             candidates=[cand(x, owner="Gita C") for x in ("hot", "due", "cold", "won")] + [cand("mine", owner="Asha K")],
             tiered=[{**tiered("hot", "B", owner="Gita C"), "opening_line": "Hi Priya, this is Gita from LawSikho — about Friday."},
                     {**tiered("due", "C", owner="Gita C"), "month_end_due": True}, tiered("cold", "C", owner="Gita C"),
                     tiered("won", "C", owner="Gita C"), tiered("mine", "A")],
             first=[{"owner": "Gita C", "lead_id": "bot", "name": "B", "source": "s", "created_ist": "x", "dials": 0,
                     "answered_calls": 0, "bot_discovery": "5 Oct", "stage": "New Lead"},
                    {"owner": "Gita C", "lead_id": "raw", "name": "R", "source": "s", "created_ist": "x", "dials": 0,
                     "answered_calls": 0, "bot_discovery": "", "stage": "New Lead"}],
             stages={"won": "Counselled lead"}, handover={"from": ["Gita C"], "to": ["Asha K", "Ravi S"]},
             contact={"due": "Asha K"})
    rows = build_rows(p)
    assert "Gita C" not in rows
    moved = {r["lead_id"]: r["owner"] for o in rows for r in rows[o] if r.get("handover_from") == "Gita C"}
    assert set(moved) == {"hot", "due", "won", "bot"} and moved["due"] == "Asha K" and moved["hot"] == "Ravi S"
    hot = next(r for r in rows["Ravi S"] if r["lead_id"] == "hot")
    assert hot["opening_line"].startswith("Hi Priya, this is Ravi, Gita's colleague, from LawSikho")
    assert hot["why"].startswith("Handover from Gita C (counsellor)")
