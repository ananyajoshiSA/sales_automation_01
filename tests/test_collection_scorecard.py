from openpyxl import load_workbook

from analytics.collection_scorecard import (callers, chances, hierarchy, next_step, priority, rows_for, why,
                                            write_workbook)
from integrations.timeutil import utc

NOW = "2026-10-24 06:00:00"
GROUPS = {"Asha Rao": ["Acefone Users", "Team Puja Malik (DSV+Women AI)", "Team Women AI"], "Himanshi": ["Elite Changemakers"],
          "Vedahi Rana": ["Elite Changemakers"], "Raunak Mandilwar": ["Elite Changemakers"],
          "New Joiner": ["Elite Changemakers"], "Someone Else": ["Group A"]}
META = {"fetched_at_utc": NOW, "team_of_owner": {"u1": "Team Puja Malik (DSV+Women AI)"},
        "team_callers": ["Asha Rao", "Vedahi Rana"], "member_groups": GROUPS}
TAG = "Independent Director Bootcamp Collection - 18th October'26"


def stage(t, cur, comment=""):
    return {"EventCode": 3002, "CreatedOn": t, "Data": [{"Key": "CurrentStage", "Value": cur}, {"Key": "Comment", "Value": comment}]}


def call(t, out=True, status="Answered", dur=200, by="Asha Rao"):
    note = f"Caller{{=}}{by}{{next}}Duration{{=}}{dur}{{next}}Status{{=}}{status}"
    return {"EventCode": 22 if out else 21, "CreatedOn": t, "ActivityFields": {"ActivityEvent_Note": note}}


def owner(t, prev, cur):
    return {"EventCode": 3001, "CreatedOn": t, "Data": [{"Key": "PreviousOwner", "Value": prev}, {"Key": "CurrentOwner", "Value": cur}]}


def lead(lid, stage_now, owner_id="u1", owner_name="Asha Rao", tag=TAG):
    return {"ProspectID": lid, "FirstName": "Lead", "LastName": lid, "OwnerId": owner_id, "OwnerIdName": owner_name,
            "ProspectStage": stage_now, "mx_Bootcamp_collections": tag}


HIST = {
    # booked 4 days ago, talked once; the lead rang back after our last dial and nobody returned it
    "fresh": [stage("2026-10-20 05:00:00", "Booking fees received"), call("2026-10-20 06:00:00"),
              call("2026-10-21 10:00:00", out=False, status="Missed", dur=0)],
    # paid 2 days after booking
    "paid": [stage("2026-10-18 05:00:00", "Booking fees received"), call("2026-10-18 06:00:00"),
             stage("2026-10-20 05:00:00", "Course Enrolled")],
    # dropped by Asha, then moved out of the team: still counts for Asha
    "moved": [stage("2026-10-18 05:00:00", "Booking fees received"), call("2026-10-18 06:00:00", status="NotAnswered", dur=0),
              stage("2026-10-21 05:00:00", "Not Interested", "stopped responding"), owner("2026-10-22 05:00:00", "Asha Rao", "Desk")],
    # booking fee paid months before this bootcamp: belongs to an earlier one
    "early": [stage("2026-05-01 05:00:00", "Booking fees received"), stage("2026-05-03 05:00:00", "Course Enrolled")],
}
LEADS = [lead("fresh", "Booking fees received"), lead("paid", "Course Enrolled"),
         lead("moved", "Not Interested", owner_id="u9", owner_name="Desk"), lead("early", "Course Enrolled")]


def test_hierarchy_maps_side_ad_and_tl():
    h = hierarchy(GROUPS)
    assert h["Asha Rao"] == ("Puja / Anmol", "Anmol Gakhar", "Nancy Bajaj")
    assert h["Vedahi Rana"] == ("Mayur", "Sahil Arora", "Sonal Tripathi")
    assert h["Raunak Mandilwar"] == ("Mayur", "Raunak Mandilwar", "(AD)")
    assert h["Himanshi"] == ("Mayur", "Raunak Mandilwar", "Ayush Srivastava")
    assert h["New Joiner"] == ("Mayur", "Not known", "Not known")
    assert "Someone Else" not in h


def test_rows_credit_the_last_team_caller_and_leave_out_early_bookings():
    rows, recon = rows_for(LEADS, HIST, META)
    r = {x["lead_id"]: x for x in rows}
    assert recon["booking fee paid for an earlier bootcamp (left out)"] == 1 and "early" not in r
    assert r["moved"]["caller"] == "Asha Rao" and r["moved"]["moved_out"] and r["moved"]["current_owner"] == "Desk"
    assert r["moved"]["side"] == "Puja / Anmol" and r["moved"]["outcome"] == "lost"
    assert r["paid"]["days_to_collect"] == 2.0 and not r["paid"]["moved_out"]
    assert r["fresh"]["call_not_returned"] and r["fresh"]["lead_name"] == "Lead fresh"


def test_pipeline_priority_reason_and_step():
    rows, _ = rows_for(LEADS, HIST, META)
    fresh = next(x for x in rows if x["lead_id"] == "fresh")
    assert priority(fresh) == "1 · Today"
    assert why(fresh, utc(NOW)).startswith("Lead called, not called back")
    assert next_step(fresh).startswith("Call back first")
    old = {**fresh, "days_since_booking": 75, "call_not_returned": False, "days_since_real": 40}
    assert priority(old) == "3 · TL decides" and next_step(old).startswith("TL: final call")
    assert next_step({**fresh, "dials": 0, "call_not_returned": False}).startswith("First call not made yet")


def test_chances_measure_later_payment_among_still_pending():
    def r(days_since, outcome, collect=None, loss=None):
        return {"days_since_booking": days_since, "outcome": outcome, "days_to_collect": collect, "days_to_loss": loss}
    rows = [r(100, "collected", 2), r(100, "collected", 20), r(100, "lost", loss=10), r(100, "open"), r(5, "collected", 1)]
    c = chances(rows)
    assert c["0–7 days"] == (50.0, 4)       # the 5-day-old lead is too young to count
    assert c["8–14 days"] == (33.3, 3)      # paid on day 2 is out; paid day 20 is in
    assert c["15–30 days"] == (50.0, 2)     # lost on day 10 is out


def test_callers_box_compares_recovery_to_bar_and_effort_to_median():
    def row(caller, collected, n, d24=True, days=5):
        return [{"mature": True, "side": "Mayur", "ad": "A", "tl": "T", "caller": caller, "dialled_24h": d24,
                 "real_convs": 1, "attempt_days": days, "callbacks_due": 0, "callbacks_missed": 0, "stalled": False,
                 "moved_out": False, "reason": "", "outcome": "collected" if i < collected else "lost"} for i in range(n)]
    rows = row("Good", 18, 20) + row("Lazy", 10, 20, d24=False, days=1) + row("Trying", 10, 20, days=9)
    c = {x["caller"]: x for x in callers(rows, 85)}
    assert c["Good"]["box"] == "Performing well" and c["Good"]["gap_to_bar"] == 0
    assert c["Lazy"]["box"] == "Effort first" and c["Lazy"]["gap_to_bar"] == 7
    assert c["Trying"]["box"] == "Effort there, not converting"


def test_workbook_has_every_sheet_and_lists_the_pending_lead(tmp_path):
    rows, recon = rows_for(LEADS, HIST, META)
    out = write_workbook(str(tmp_path / "s.xlsx"), rows, recon, META, 85)
    wb = load_workbook(tmp_path / "s.xlsx")
    assert wb.sheetnames == ["Summary", "Recovery pipeline", "Second chance", "Callers", "Weekly cohorts", "Bootcamps"]
    pipe = wb["Recovery pipeline"]
    assert pipe["A3"].value == "1 · Today" and pipe["H3"].value == "fresh" and pipe["A4"].value is None
    assert out["pipeline"] == 1 and out["second_chance"] == 1
    assert wb["Second chance"]["D3"].value == "Asha Rao" and wb["Second chance"]["E3"].value == "Desk"
    assert "24 Oct 2026, 11:30 IST" in wb["Summary"]["A1"].value
