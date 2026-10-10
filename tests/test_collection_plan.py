from datetime import date

from openpyxl import load_workbook

from analytics import collection_plan as cp


def row(lid, caller="Asha Rao", team="Elite Changemakers", outlook="Workable with follow-up", blocker="No clear reason recorded",
        days_open="5", flags="", chance="0.2", kind="Bootcamp"):
    return {"lead_id": lid, "kind": kind, "team": team, "caller": caller, "course": "Independent Director",
            "booked_ist": "2026-10-05 12:00", "days_open": days_open, "stage": "Call Back Later", "flags": flags,
            "outlook": outlook, "blocker": blocker, "situation": "s", "next_action": "n", "what_to_say": "w",
            "process_issues": "No dated next step; Lead's call not returned", "chance_3d": chance, "chance_3d_low": "0.1",
            "chance_3d_high": "0.3", "chance_14d": "0.4"}


def test_tier_of_follows_the_reading_and_flags():
    assert cp.tier_of(row("a", outlook="Check payment first", flags="Lead called and was not called back")) == "P"
    assert cp.tier_of(row("a", outlook="Likely lost")) == "X"
    assert cp.tier_of(row("a", outlook="Ready to close", flags="Lead called and was not called back the same day")) == "M"
    assert cp.tier_of(row("a", outlook="Ready to close")) == "A"
    assert cp.tier_of(row("a", blocker="Loan or EMI paperwork in progress")) == "L"
    assert cp.tier_of(row("a", outlook="Long shot")) == "R"
    assert cp.tier_of(row("a", days_open="14")) == "B" and cp.tier_of(row("a", days_open="15")) == "C"
    assert cp.clean_name("Malti Jain Jain") == "Malti Jain"


def test_plan_rows_drop_closed_leads_order_by_group_and_spread_older_leads():
    rows = [row("old1", days_open="20", chance="0.05"), row("old2", days_open="25", chance="0.04"),
            row("ready", outlook="Ready to close", chance="0.3"), row("paid", outlook="Ready to close")]
    contacts = {"old1": {"name": "Ravi Das Das", "phone": "+91-1"}, "paid": {"stage": "Collections done"}}
    cp.DAY1_LOAD, cp.DAY2_BACKLOG = 2, 1          # one older lead fits day 1, the next goes to day 2
    try:
        teams, gone = cp.plan_rows(rows, contacts, date(2026, 10, 11))
    finally:
        cp.DAY1_LOAD, cp.DAY2_BACKLOG = 30, 10
    assert gone == {"Collections done": 1}
    plan = teams["Elite Changemakers"]["Asha Rao"]
    assert [r["lead_id"] for r in plan] == ["ready", "old1", "old2"]
    assert plan[1]["name"] == "Ravi Das" and plan[1]["phone"] == "+91-1"
    assert plan[1]["first_day"] == "Sun 11 Oct" and plan[2]["first_day"] == "Mon 12 Oct"
    assert plan[0]["when"].startswith("Sun 11 Oct 11:00")


def test_workbook_and_html(tmp_path):
    teams, _ = cp.plan_rows([row("a", outlook="Ready to close"), row("b", caller="Ravi Das", outlook="Check payment first"),
                             row("c", outlook="Long shot", kind="Community")], {}, date(2026, 10, 11))
    callers = teams["Elite Changemakers"]
    days = cp.day_names(date(2026, 10, 11))
    path = tmp_path / "plan.xlsx"
    cp.write_workbook(str(path), "Elite Changemakers", callers, days, "2026-10-10 16:17")
    wb = load_workbook(path)
    assert wb.sheetnames == ["How to use", "Team summary", "Team leader actions", "Asha Rao", "Ravi Das"]
    ws = wb["Asha Rao"]
    assert ws.cell(5, cp.COL["Call group"]).value == "A · Ready to close"
    assert ws.cell(6, cp.COL["Call group"]).value == "R · Long shot"
    assert ws.cell(5, cp.COL["Status"]).value.startswith('=IF(')
    assert wb["Team leader actions"].cell(4, 1).value == "P · Check payment first"
    report = {"by_kind_team": [{"kind": "Bootcamp", "team": "Elite Changemakers", "pool": 10, "collected": 6, "collected_%_now": 60.0,
                                "lost": 1, "open": 3, "expected_next_3d": 0.6, "range_3d": "0-2", "expected_next_45d": 1.0,
                                "range_45d": "0-3", "projected_%_45d": 70.0, "needed_for_75%": 2, "needed_for_75%_share_of_open": "67%",
                                "needed_for_80%": 2, "needed_for_80%_share_of_open": "67%"}],
              "history": [{"key": "Bootcamp / 0-3 days / Loan pending", "snapshots": 40, "collected_14d_%": 80.0},
                          {"key": "Bootcamp / 0-3 days / Call back later", "snapshots": 40, "collected_14d_%": 50.0}]}
    page = cp.build_html("Elite Changemakers", callers, report, date(2026, 10, 11), "2026-10-10 16:17", 0)
    assert "Sun 11 to Tue 13 October 2026" in page and "for Mayur Sachdeva" in page
    assert "50–80% of bootcamp leads" in page and "not enough history" in page
    assert page.count('<section class="caller">') == 2
