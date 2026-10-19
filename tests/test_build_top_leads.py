import json

from openpyxl import load_workbook

from scripts.build_top_leads import main, top_leads


def r(lid, tier, prob, **kw):
    return {"lead_id": lid, "name": f"Lead {lid}", "tier": tier, "prob_3d": prob, "why": "w", "next_action": "n", **kw}


PLAN = {"Asha K": [r("b1", "B", 20), r("a1", "A", 30), r("c1", "C", 5), r("h1", "A", 30, handover_from="Gita C")],
        "Ravi S": [r("a2", "A", 30, month_end_due=True), r("b2", "B", 20, rank=9)], "Lead Er": [r("x", "A", 50)]}


def test_ranked_by_chance_then_promise_with_owner_now():
    rows = top_leads(PLAN, {"h1": "Gita C", "a1": "Asha K"}, n=5, leader="Lead Er")
    assert [x["lead_id"] for x in rows] == ["a2", "a1", "h1", "b2", "b1"]
    h1 = next(x for x in rows if x["lead_id"] == "h1")
    assert (h1["owner_now"], h1["call_by"]) == ("Gita C", "Asha K")


def test_files(tmp_path, monkeypatch):
    (tmp_path / "plan.json").write_text(json.dumps(PLAN))
    (tmp_path / "snap.json").write_text(json.dumps({"leads": [{"ProspectID": "h1", "OwnerIdName": "Gita C"}]}))
    (tmp_path / "meta.json").write_text(json.dumps({"team": "Team Test", "leader": "Lead Er", "counsellors": ["Gita C"]}))
    out = str(tmp_path / "top")
    monkeypatch.setattr("sys.argv", ["x", str(tmp_path / "plan.json"), str(tmp_path / "snap.json"), out,
                                     "--meta", str(tmp_path / "meta.json"), "--top", "3"])
    main()
    ws = load_workbook(out + ".xlsx")["Top 3"]
    assert [ws.cell(row=k, column=20).value for k in (5, 6, 7)] == ["a2", "a1", "h1"]
    assert [ws.cell(row=7, column=c).value for c in (4, 5, 6)] == ["Gita C", "Counsellor", "Asha K"]
    assert "Owner now: Gita C – change it" in open(out + ".html").read()
