import json

from openpyxl import load_workbook

from scripts.build_focus_list import focus_rows, main
from scripts.build_plan_pdf import build


def r(lid, tier, prob, **kw):
    return {"lead_id": lid, "name": f"Lead {lid}", "tier": tier, "prob_3d": prob, "why": "w", "next_action": "n", **kw}


PLAN = {"Asha K": [r("a1", "A", 25), r("b1", "B", 18), r("b2", "B", 10), r("b3", "B", 12, month_end_due=True), r("f1", "F", 5)],
        "Ravi S": [r("a2", "A", 40), r("b4", "B", 8, month_end_due=True)], "Lead Er": [r("a3", "A", 50)]}


def test_focus_keeps_a_strong_b_and_due_promises_in_order():
    rows = focus_rows(PLAN, leader="Lead Er")
    assert [x["lead_id"] for x in rows] == ["a2", "a1", "b3", "b1"]
    assert rows[0]["owner"] == "Ravi S"


def test_focus_files(tmp_path, monkeypatch):
    (tmp_path / "plan.json").write_text(json.dumps(PLAN))
    (tmp_path / "meta.json").write_text(json.dumps({"team": "Team Test", "leader": "Lead Er", "day": "Monday", "next_day": "Tuesday"}))
    out = str(tmp_path / "focus")
    monkeypatch.setattr("sys.argv", ["x", str(tmp_path / "plan.json"), out, "--meta", str(tmp_path / "meta.json")])
    main()
    ws = load_workbook(out + ".xlsx")["Focus list"]
    assert load_workbook(out + ".xlsx").sheetnames == ["Focus list", "Focus by caller"]
    assert [ws.cell(row=k, column=20).value for k in range(5, 9)] == ["a2", "a1", "b3", "b1"]
    assert ws.cell(row=4, column=17).value == "Monday outcome"
    doc = open(out + ".html").read()
    assert "Where Monday's and Tuesday's enrolments can come from" in doc and "Lead a3" not in doc


def test_one_consolidated_workbook_and_pdf(tmp_path, monkeypatch):
    from openpyxl import Workbook
    plan_wb = Workbook()
    plan_wb.active.title = "How to use"
    plan_wb.create_sheet("Team summary")
    plan_wb.save(tmp_path / "plan.xlsx")
    (tmp_path / "plan.json").write_text(json.dumps(PLAN))
    meta = {"team": "Team Test", "leader": "Lead Er", "day": "Monday", "next_day": "Tuesday", "focus": True}
    (tmp_path / "meta.json").write_text(json.dumps(meta))
    out = str(tmp_path / "all")
    monkeypatch.setattr("sys.argv", ["x", str(tmp_path / "plan.json"), out, "--meta", str(tmp_path / "meta.json"),
                                     "--into", str(tmp_path / "plan.xlsx")])
    main()
    assert load_workbook(out + ".xlsx").sheetnames == ["How to use", "Focus list", "Focus by caller", "Team summary"]
    build(str(tmp_path / "plan.json"), out + "_plan.html", str(tmp_path / "meta.json"))
    doc = open(out + "_plan.html").read()
    assert doc.index("Start here: where Monday's and Tuesday's") < doc.index("1. The target")
