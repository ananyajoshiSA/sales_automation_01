import json
from datetime import datetime

from analytics.daily_plan import content, mailer, rows, state, status, transcripts, workbook
from analytics.daily_plan.common import Snap, previous_working_day
from analytics.daily_plan.stats import day_stats

USERS = [{"ID": "u1", "FirstName": "Asha", "LastName": "K"}, {"ID": "u2", "FirstName": "Ravi", "LastName": "S"},
         {"ID": "u0", "FirstName": "Tara", "LastName": "L"}]
LEADS = [{"ProspectID": f"L{i}", "FirstName": f"Lead{i}", "Phone": f"+91-90000000{i:02d}", "OwnerIdName": "Asha K" if i % 2 else "Ravi S",
          "ProspectStage": "Follow Up For Closure", "CreatedOn": "2026-09-01 05:00:00"} for i in range(1, 9)]


def call(user, lead, utc, status="NotAnswered", dur=0, direction="outbound"):
    return {"activity_id": f"{user}{lead}{utc}{direction}", "user_id": user, "lead_id": lead, "start_utc": utc, "status": status,
            "duration": dur, "direction": direction, "lead_number": "90000000" + lead[1:].rjust(2, "0"), "display_number": "918065990718"}


def snap(calls, fetched="2026-10-10T08:30:00+00:00", leads=LEADS):
    return Snap({"users": USERS, "fetched_at": fetched, "leads": leads, "calls": calls, "zip_activities": []}, leader="Tara L")


def test_state_keeps_no_lead_ids_and_resolves_back():
    by_owner = {"Asha K": [{"lead_id": "L1", "tier": "A", "chance": 30}], "Ravi S": [{"lead_id": "L2", "tier": "B", "chance": 10, "verify": True}]}
    st = state.to_state("2026-10-10", by_owner, [{"lead_id": "L1", "group": 1, "check_by": "11:30"}])
    assert "L1" not in json.dumps(st) and "Lead1" not in json.dumps(st)
    got = state.resolve(st, ["L1", "L2", "L3"])
    assert got["L1"]["group"] == 1 and got["L2"]["verify"] and "L3" not in got


def test_day_stats_measures_the_a_block_and_missed_calls():
    sheet = {"L1": {"owner": "Asha K", "tier": "A"}, "L3": {"owner": "Asha K", "tier": "A"}, "L2": {"owner": "Ravi S", "tier": "B"}}
    cs = [call("u1", "L1", "2026-10-09 05:00:00"), call("u1", "L1", "2026-10-09 05:10:00", "Answered", 300),  # redial in 10 min, reached
          call("u1", "L5", "2026-10-09 06:00:00"),                                                               # off-sheet dial
          call("u0", "L2", "2026-10-09 07:00:00", "Missed", 0, "inbound")]                                        # Ravi's lead rang, unreturned
    st = day_stats(snap(cs), "2026-10-09", sheet)
    a, r, t = st["callers"]["Asha K"], st["callers"]["Ravi S"], st["team"]
    assert (a["dials"], a["real"], a["on_sheet_pct"], a["a_tried"], a["a_twice"], a["a_reached"]) == (3, 1, 67, 1, 1, 1)
    assert (t["a_total"], t["a_tried"], t["a_first_unanswered"], t["a_redial15"]) == (2, 1, 1, 1)
    assert t["a_not_tried"] == ["L3"] and t["absent"] == ["Ravi S"]
    assert r["missed_in"] == 1 and r["missed_unreturned"] == 1


def test_previous_working_day_skips_a_quiet_sunday():
    cs = [call("u1", "L1", f"2026-10-10 05:{m:02d}:00") for m in range(55)]
    s = snap(cs)
    assert previous_working_day(s, "2026-10-12") == "2026-10-10"


def test_enrolled_or_paid_leads_never_go_back_on_a_sheet():
    reads = {"L1": {"lead_id": "L1", "name": "Lead1", "status": "paid_new", "tier": "D", "chance": 0, "payment_evidence": "Rs 3,000 paid on the call"},
             "L3": {"lead_id": "L3", "name": "Lead3", "status": "21day", "tier": "D", "chance": 0},
             "L4": {"lead_id": "L4", "name": "Lead4", "status": "active", "tier": "B", "chance": 12},
             "L5": {"lead_id": "L5", "name": "Lead5", "status": "active", "tier": "A", "chance": 40}}
    leads = [dict(l, ProspectStage="Course Enrolled") if l["ProspectID"] in ("L5", "L6") else l for l in LEADS]
    cs = [call("u0", "L4", "2026-10-09 10:00:00", "Missed", 0, "inbound")]
    prev = {"L7": {"owner": "Asha K", "tier": "C", "chance": 3}, "L6": {"owner": "Ravi S", "tier": "A", "chance": 30},
            "L8": {"owner": "Ravi S", "tier": "P", "chance": 0}}
    R = rows.build_rows(snap(cs, leads=leads), "2026-10-10", reads, prev, [])
    flat = {r["lead_id"]: r for rs in R["by_owner"].values() for r in rs}
    assert not {"L1", "L5", "L6", "L8"} & set(flat)                       # paid on a call, Course Enrolled, or an old paid row
    assert sorted(p["lead_id"] for p in R["paid"]) == ["L1", "L5"]
    assert "L3" not in flat and R["dropped"][0]["status"] == "21day"
    assert flat["L4"]["tier"] == "M" and flat["L4"]["missed"] == 1
    assert flat["L7"]["source"] == "carried" and flat["L7"]["tier"] == "C"


def test_leads_marked_not_interested_or_later_stay_off_the_sheets_unless_they_ring_us():
    stage = {"L1": "Not Interested", "L2": "May buy later", "L3": "Irrelevant lead", "L4": "May buy later", "L5": "May buy later", "L6": "Not Interested"}
    leads = [dict(l, ProspectStage=stage.get(l["ProspectID"], l["ProspectStage"])) for l in LEADS]
    reads = {lid: {"lead_id": lid, "name": lid, "status": "active", "tier": t, "chance": 5}
             for lid, t in (("L1", "C"), ("L2", "B"), ("L4", "C"), ("L5", "A"))}
    prev = {"L3": {"owner": "Asha K", "tier": "C", "chance": 2}, "L6": {"owner": "Ravi S", "tier": "B", "chance": 8}}
    cs = [call("u0", "L4", "2026-10-10 03:00:00", "Missed", 0, "inbound")]                # L4 rang us and is still waiting
    R = rows.build_rows(snap(cs, leads=leads), "2026-10-10", reads, prev, [])
    flat = {r["lead_id"] for rs in R["by_owner"].values() for r in rs}
    assert not {"L1", "L2", "L3", "L6"} & flat                                             # caller's disposition respected
    assert {"L4", "L5"} <= flat                                                            # rang us since / about to pay
    assert sorted(d["lead_id"] for d in R["dropped"] if d["status"] == "shelved") == ["L1", "L2"]


def test_rs10_bootcamp_registration_is_a_prospect_not_an_enrollment():
    from analytics.daily_plan import review
    leads = [dict(l, ProspectStage="Course Enrolled") if l["ProspectID"] in ("L1", "L3") else l for l in LEADS]
    reads = {"L1": {"lead_id": "L1", "name": "Lead1", "status": "bootcamp", "tier": "B", "chance": 8, "payment_evidence": "Rs 10 bootcamp registration"},
             "L3": {"lead_id": "L3", "name": "Lead3", "status": "paid_new", "tier": "D", "chance": 0}}
    enrolled = [{"ProspectID": lid, "FirstName": f"Lead{lid[1:]}", "OwnerIdName": "Asha K", "first_enrolled": "2026-10-09 06:00:00", "history": []}
                for lid in ("L1", "L3")]
    s = snap([], leads=leads)
    R = rows.build_rows(s, "2026-10-10", reads, {}, enrolled)
    flat = {r["lead_id"]: r for rs in R["by_owner"].values() for r in rs}
    assert flat["L1"]["tier"] == "B" and "Rs 10 bootcamp" in flat["L1"]["why"] and "L3" not in flat
    assert [e["lead_id"] for e in review.enrollments(s, "2026-10-09", enrolled, reads)] == ["L3"]
    assert [b["name"] for b in review.bootcamp_registrations(s, "2026-10-09", enrolled, reads)] == ["Lead1"]
    st = state.to_state("2026-10-10", R["by_owner"], [{"lead_id": "L1", "group": 3, "check_by": "15:30"}])
    S = status.build(s, "2026-10-10", state.resolve(st, s.leads), [], "Tara L")
    assert [p["lead_id"] for p in S["priority"]] == ["L1"]


def test_an_old_payment_is_a_stage_fix_not_the_days_enrollment():
    from analytics.daily_plan import review
    reads = {"L1": {"lead_id": "L1", "name": "Lead1", "status": "paid_new", "paid_on": "2026-08-13", "payment_evidence": "Rs 3,000 in Aug"},
             "L3": {"lead_id": "L3", "name": "Lead3", "status": "paid_new", "paid_on": "2026-10-09", "payment_evidence": "Rs 3,000 on the call"},
             "L4": {"lead_id": "L4", "name": "Lead4", "status": "paid_new", "payment_evidence": "paid, date unknown"}}
    s = snap([])
    assert [e["lead_id"] for e in review.enrollments(s, "2026-10-09", [], reads)] == ["L3"]
    assert sorted(x["lead_id"] for x in review.paid_unstaged(s, "2026-10-09", reads)) == ["L1", "L4"]
    plan = {"L1": {"owner": "Asha K", "group": 1}, "L3": {"owner": "Asha K", "group": 1}}
    res = {o["lead_id"]: o["result"] for o in review.priority_outcomes(s, "2026-10-09", plan, reads, {"L3"})}
    assert res == {"L1": "Paid earlier", "L3": "Won"}


def test_priority_groups_respect_the_day_the_lead_asked_for():
    R = {"today": "2026-10-10", "by_owner": {"Asha K": [
        {"lead_id": "a", "name": "A", "tier": "A", "chance": 30, "callback_requested": ""},
        {"lead_id": "b", "name": "B", "tier": "B", "chance": 8, "callback_requested": "Sat 10 Oct 18:00"},
        {"lead_id": "c", "name": "C", "tier": "B", "chance": 12, "callback_requested": "Monday evening"},
        {"lead_id": "d", "name": "D", "tier": "A", "chance": 40, "callback_requested": "", "exact_ask": "Send the payment link today"},
        {"lead_id": "f", "name": "F", "tier": "B", "chance": 12, "callback_requested": "Monday 12 Oct (on a train Saturday; do not call today)"},
        {"lead_id": "e", "name": "E", "tier": "C", "chance": 2, "callback_requested": ""}]}}
    g = {p["lead_id"]: (p["group"], p["check_by"]) for p in content.priority(R, "2026-10-10")}
    assert g == {"d": (0, "12:00"), "a": (1, "11:30"), "b": (2, "18:00"), "c": (5, "later"), "f": (5, "later")}


def test_workbook_counts_paid_without_asking_for_proof(tmp_path):
    R = {"by_owner": {"Asha K": [{"lead_id": "L1", "name": "Lead1", "tier": "A", "chance": 30, "phone": "9000000001"}]}}
    C = {"leader": "Tara L", "review": {"day_label": "Fri 9 Oct"}, "slots": content.SLOTS, "shivangi_title": "t", "shivangi_intro": "i",
         "priority": [], "priority_groups": content.GROUPS, "checks": [], "feedback": [], "changes": [], "caller_note": "n",
         "owner_order": ["Asha K"], "targets": {}, "how_to": [], "day_label": "Sat"}
    out = workbook.write(R, C, str(tmp_path / "p.xlsx"))
    from openpyxl import load_workbook
    ws = load_workbook(out)["Asha K"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=IF(")]
    assert formulas and all('"Enrolled"' in f and "proof" not in f and '"Already a student"' in f for f in formulas)


def test_status_lists_untouched_tried_and_unreturned():
    plan = {"L1": {"owner": "Asha K", "tier": "A", "group": 1, "check_by": "11:30"},
            "L3": {"owner": "Asha K", "tier": "A", "group": 1, "check_by": "11:30"},
            "L5": {"owner": "Asha K", "tier": "B", "group": 2, "check_by": "18:00"},
            "L7": {"owner": "Asha K", "tier": "A", "group": 1, "check_by": "11:30"},
            "L8": {"owner": "Ravi S", "tier": "B", "group": 5, "check_by": "later"}}                   # asked for Monday
    cs = [call("u1", "L3", "2026-10-10 05:00:00"), call("u0", "L2", "2026-10-10 06:00:00", "Missed", 0, "inbound")]
    leads = [dict(l, ProspectStage="Course Enrolled") if l["ProspectID"] == "L7" else l for l in LEADS]
    S = status.build(snap(cs, leads=leads), "2026-10-10", plan, [], "Tara L")
    by = {p["lead_id"]: p["status"] for p in S["priority"]}
    assert by == {"L1": "untouched", "L3": "tried", "L5": "untouched"}
    text = " ".join(S["todo"])
    assert "Ravi S" in text and "Lead1" in text and "Lead2" in text and "Lead5 18:00" in text
    assert "9000000002" not in status.text_summary(S, "Team")


def test_mailer_needs_settings_and_sends_attachments(tmp_path, monkeypatch):
    for k in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "MAIL_RELAY_URL", "MAIL_RELAY_KEY"):
        monkeypatch.delenv(k, raising=False)
    assert mailer.send("s", "<p>h</p>", "t", ["a@x.in"]) is False
    monkeypatch.setenv("SMTP_HOST", "smtp.x.in"); monkeypatch.setenv("SMTP_USER", "r@x.in"); monkeypatch.setenv("SMTP_PASSWORD", "p")
    f = tmp_path / "plan.pdf"; f.write_bytes(b"%PDF")
    sent = []

    class FakeSMTP:
        def __init__(self, host, port, timeout): sent.append((host, port))
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def starttls(self): pass
        def login(self, u, p): pass
        def send_message(self, m): sent.append(m)

    assert mailer.send("Plan", "<p>h</p>", "t", ["a@x.in", "b@x.in"], [str(f)], smtp_factory=FakeSMTP)
    msg = sent[1]
    assert sent[0] == ("smtp.x.in", 587) and msg["To"] == "a@x.in, b@x.in"
    assert [p.get_filename() for p in msg.iter_attachments()] == ["plan.pdf"]


def test_transcript_runs_stay_within_limits_and_resume(tmp_path):
    made = []

    class FakeClient:
        requests_made = 0
        def search_raw(self, nums):
            assert len(nums) <= 90
            made.append(len(nums)); self.requests_made = -(-len(nums) // 10)
            return {n: {"sales_call": []} for n in nums}

    pauses = []
    nums = [f"9190000{i:05d}" for i in range(200)]
    transcripts.fetch(nums, str(tmp_path), pause=5, client_factory=FakeClient, sleep=pauses.append)
    assert made == [90, 90, 20] and pauses == [5, 5]
    transcripts.fetch(nums, str(tmp_path), pause=5, client_factory=FakeClient, sleep=pauses.append)
    assert made == [90, 90, 20]


def test_pipeline_runs_on_python_311():
    """Scheduled cloud sessions start on python3 = 3.11, which rejects 3.12-only f-string quoting."""
    import pathlib
    import shutil
    import subprocess

    import pytest

    py = shutil.which("python3.11")
    if not py:
        pytest.skip("python3.11 not installed")
    files = [str(p) for p in pathlib.Path("analytics/daily_plan").glob("*.py")]
    r = subprocess.run([py, "-m", "py_compile", *files], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr


def test_pipeline_horizon_candidates_and_state(tmp_path):
    from analytics.daily_plan import pipeline
    assert pipeline.horizon("2026-10-10") == "2026-10-31" and pipeline.horizon("2026-10-29") == "2026-11-30"
    leads = [dict(l, ProspectStage="Course Enrolled") if l["ProspectID"] == "L4" else l for l in LEADS]
    cs = [call("u1", "L1", "2026-10-10 06:00:00", "Answered", 300), call("u1", "L2", "2026-10-10 06:10:00", "Answered", 40)]
    s = snap(cs, leads=leads)
    last_night = {"L3": {"month_chance": 20, "stage": "2"}, "L4": {"month_chance": 30, "stage": "1"}}
    plan = {"L5": {"tier": "A", "group": 1}, "L6": {"tier": "C", "group": None}}
    assert pipeline.candidates(s, "2026-10-10", last_night, plan) == {"L1", "L3", "L5"}  # L4 enrolled, L2 too short, L6 nurture
    reads = [{"lead_id": "L1", "name": "Lead1", "owner": "Asha K", "include": True, "month_chance": 12, "stage_reached": "3 Fee and plan discussed",
              "expected_window": "12–14 Oct"},
             {"lead_id": "L3", "name": "Lead3", "owner": "Asha K", "include": True, "month_chance": 30, "stage_reached": "2 Agreed, one blocker",
              "expected_window": "by 28 Oct"},
             {"lead_id": "L5", "name": "Lead5", "owner": "Ravi S", "include": False, "month_chance": 0, "exclude_reason": "already enrolled"}]
    inc = pipeline.included(reads)
    assert [r["lead_id"] for r in inc] == ["L3", "L1"]
    S = pipeline.summary(inc, "2026-10-10")
    assert round(S["expected_low"], 2) == 0.42 and S["weeks"][0][0] == 1 and S["weeks"][2][0] == 1
    st = pipeline.to_state("2026-10-10", inc)
    assert "L3" not in json.dumps(st) and "Lead3" not in json.dumps(st)
    state.save(st, str(tmp_path), "pipeline_")
    assert state.latest_before(str(tmp_path), "2026-10-11", "pipeline_")["date"] == "2026-10-10"
    assert state.latest_before(str(tmp_path), "2026-10-10", "pipeline_") is None
    text = pipeline.brief("by {horizon_short}; {today_long}; {days_left} days", "2026-10-10", "2026-10-31")
    assert text == "by 31 Oct; Saturday 10 Oct 2026; 21 days"


def test_fresh_leads_allocated_today_and_what_they_produced():
    from analytics.daily_plan import pipeline
    leads = [dict(l) for l in LEADS]
    leads[0].update(mx_Assigned_On="2026-10-10 04:00:00", CreatedOn="2026-10-09 10:00:00")   # L1 new, Asha
    leads[1].update(mx_Assigned_On="2026-10-10 04:00:00")                                      # L2 old lead reassigned, Ravi
    leads[2].update(mx_Assigned_On="2026-10-10 05:00:00", mx_Next_follow_up_date="2026-10-12 06:00:00")  # L3 reached, LSQ follow-up
    leads[3].update(mx_Assigned_On="2026-10-01 05:00:00")                                      # L4 not today
    cs = [call("u1", "L1", "2026-10-10 06:00:00", "Answered", 400), call("u2", "L2", "2026-10-10 06:00:00"),
          call("u1", "L3", "2026-10-10 07:00:00", "Answered", 70)]
    s = snap(cs, leads=leads)
    alloc = pipeline.fresh(s, "2026-10-10")
    assert alloc == {"L1": {"owner": "Asha K", "new": True}, "L2": {"owner": "Ravi S", "new": False}, "L3": {"owner": "Asha K", "new": False}}
    assert pipeline.candidates(s, "2026-10-10", {}, {}, alloc) == {"L1", "L3"}
    reads = [{"lead_id": "L1", "name": "Lead1", "owner": "Asha K", "include": True, "month_chance": 12, "stage_reached": "3 Fee and plan discussed",
              "followup_agreed": True, "followup_when": "Mon 12 Oct 18:00"}]
    F = pipeline.fresh_outcomes(s, "2026-10-10", alloc, reads)
    a, r = F["callers"]["Asha K"], F["callers"]["Ravi S"]
    assert (a["allocated"], a["new"], a["dialled"], a["reached"], a["real"], a["followup"], a["pipeline"]) == (2, 1, 2, 2, 1, 2, 1)
    assert (r["allocated"], r["reassigned"], r["dialled"], r["reached"], r["followup"]) == (1, 1, 1, 0, 0)
    assert [g["lead_id"] for g in F["good"]] == ["L1", "L3"] and F["good"][0]["followup"] == "Mon 12 Oct 18:00"
    assert "Fresh leads allocated Sat 10 Oct" in pipeline.fresh_html(F, "Sat 10 Oct")


class _FakeLSQ:
    """Just enough of LeadSquaredClient for the incremental fetches."""

    def __init__(self, acts=None, enrolled=None):
        self.acts, self.enrolled, self.reads, self.histories = acts or {}, enrolled or [], [], []

    def get_users_in_group(self, group):
        return [{"ID": "u1", "FirstName": "Asha", "LastName": "K"}]

    def iter_leads_by_group(self, group, columns=None):
        return iter([{"ProspectID": "L1"}, {"ProspectID": "L2"}])

    def iter_activities_started(self, ev, d0, d1, edited_since=None):
        self.reads.append((ev, edited_since))
        return iter(self.acts.get(ev, []))

    def iter_leads(self, *a, **k):
        return iter(self.enrolled)

    def get_lead_activities(self, lid, **k):
        self.histories.append(lid)
        return {"ProspectActivities": []}


def _act(aid, lead, created, modified):
    return {"ProspectActivityId": aid, "RelatedProspectId": lead, "ActivityEvent": 22, "CreatedOn": created, "ModifiedOn": modified,
            "Owner": "u1", "ActivityEvent_Note": ""}


def test_incremental_snapshot_reuses_old_rows_and_takes_fresh_edits(monkeypatch, tmp_path):
    import scripts.fetch_team_data as ftd
    from integrations.timeutil import utc as _utc

    prev = {"group": "Team X", "start": "2026-10-07", "end": "2026-10-10", "fetched_at": "2026-10-10T16:00:00+00:00",
            "calls": [{"activity_id": "a1", "lead_id": "L1", "user_id": "u1", "start_utc": "2026-10-08 05:00:00", "modified_utc": "old"},
                      {"activity_id": "a2", "lead_id": "L1", "user_id": "u1", "start_utc": "2026-10-10 15:00:00", "modified_utc": "old"},
                      {"activity_id": "gone", "lead_id": "L1", "user_id": "u1", "start_utc": "2026-10-06 05:00:00", "modified_utc": "old"}],
            "payments": [], "zip_activities": []}
    fake = _FakeLSQ({22: [_act("a2", "L1", "2026-10-10 15:00:00", "2026-10-10 15:05:00"), _act("a3", "L2", "2026-10-11 04:00:00", "2026-10-11 04:01:00")]})
    monkeypatch.setattr(ftd, "LeadSquaredClient", lambda: fake)
    monkeypatch.setattr(ftd, "now_utc", lambda: _utc("2026-10-11 05:00:00"))
    out = tmp_path / "snap.json"
    ftd.main("Team X", "2026-10-07", "2026-10-11", str(out), prev)
    got = {c["activity_id"]: c["modified_utc"] for c in json.load(open(out))["calls"]}
    assert got == {"a1": "old", "a2": "2026-10-10 15:05:00", "a3": "2026-10-11 04:01:00"}                 # a2's later edit wins; the out-of-window row is dropped
    assert {s for _, s in fake.reads} == {_utc("2026-10-10 14:00:00")}   # edits read from 2 h before the last pull
    assert ftd.reusable(prev, "Team X", "2026-10-06") is None         # the old pull doesn't cover the new start
    assert ftd.reusable(prev, "Team Y", "2026-10-08") is None


def test_snapshot_falls_back_to_per_lead_reads_when_bulk_is_refused(monkeypatch, tmp_path):
    import scripts.fetch_team_data as ftd
    from analytics.daily_plan import revival
    from integrations.leadsquared import LeadSquaredError
    from integrations.timeutil import utc as _utc

    class Refused(_FakeLSQ):
        def iter_leads_by_group(self, group, columns=None):
            return iter([{"ProspectID": "L1", "ProspectActivityDate_Max": "2026-10-13 06:00:00"},
                         {"ProspectID": "L2", "ProspectActivityDate_Max": "2026-10-01 06:00:00", "ModifiedOn": "2026-10-02 06:00:00"}])

        def iter_activities_started(self, *a, **k):
            self.reads.append(a[0])
            raise LeadSquaredError("Access Denied", status_code=403)

        def get_lead_activities(self, lid, **k):
            self.histories.append(lid)
            row = lambda aid, ev, t: {"Id": aid, "EventCode": ev, "CreatedOn": t, "ModifiedOn": t,
                                      "ActivityFields": {"Owner": "u1", "Status": "Answered", "Duration": "130"}}
            return {"ProspectActivities": [row("today", 22, "2026-10-13 05:00:00"), row("old", 22, "2026-10-12 05:00:00"),
                                           row("zip", 237, "2026-10-13 05:03:00"), row("email", 1, "2026-10-13 05:00:00")]}

    fake = Refused()
    monkeypatch.setattr(ftd, "LeadSquaredClient", lambda: fake)
    monkeypatch.setattr(ftd, "now_utc", lambda: _utc("2026-10-13 08:00:00"))
    monkeypatch.setattr(revival, "RATE", 1000)
    waits = []
    monkeypatch.setattr(ftd, "sleep", waits.append)
    out = tmp_path / "snap.json"
    ftd.main("Team X", "2026-10-13", "2026-10-13", str(out))
    snap = json.load(open(out))
    assert len(fake.reads) == 3 and waits == [120, 120]                       # the bulk read is tried three times first
    assert fake.histories == ["L1"]                                         # only the lead active since the start is read
    assert [(c["activity_id"], c["lead_id"], c["caller"]) for c in snap["calls"]] == [("today", "L1", "Asha K")]
    assert [z["ProspectActivityId"] for z in snap["zip_activities"]] == ["zip"]
    assert snap["source"] == "per-lead fallback" and not (tmp_path / "snap.json.hist.jsonl").exists()


def test_snapshot_retries_a_refused_bulk_read_before_falling_back(monkeypatch, tmp_path):
    import scripts.fetch_team_data as ftd
    import pytest
    from integrations.leadsquared import LeadSquaredError
    from integrations.timeutil import utc as _utc

    class RefusedOnce(_FakeLSQ):
        refusals = 1

        def iter_activities_started(self, ev, d0, d1, edited_since=None):
            if self.refusals:
                self.refusals -= 1
                raise LeadSquaredError("Access Denied", status_code=403)
            return super().iter_activities_started(ev, d0, d1, edited_since)

    fake = RefusedOnce({22: [_act("a1", "L1", "2026-10-13 05:00:00", "2026-10-13 05:01:00")]})
    monkeypatch.setattr(ftd, "LeadSquaredClient", lambda: fake)
    monkeypatch.setattr(ftd, "now_utc", lambda: _utc("2026-10-13 08:00:00"))
    waits = []
    monkeypatch.setattr(ftd, "sleep", waits.append)
    out = tmp_path / "snap.json"
    ftd.main("Team X", "2026-10-13", "2026-10-13", str(out))
    snap = json.load(open(out))
    assert waits == [120] and fake.histories == []                            # one pause, then the bulk read worked
    assert snap["source"] == "bulk" and [c["activity_id"] for c in snap["calls"]] == ["a1"]

    class Broken(_FakeLSQ):
        def iter_activities_started(self, *a, **k):
            raise LeadSquaredError("Server error", status_code=500)

    monkeypatch.setattr(ftd, "LeadSquaredClient", lambda: Broken())
    waits.clear()
    with pytest.raises(LeadSquaredError):                                     # only a 403 is waited out
        ftd.main("Team X", "2026-10-13", "2026-10-13", str(out))
    assert waits == []


def test_enrolled_reuses_known_first_enrolments(tmp_path):
    from analytics.daily_plan import fetch

    leads = [{"ProspectID": p, "ModifiedOn": "2026-10-11 04:00:00"} for p in ("L1", "L2", "L3")]
    cache = {"L1": {"ProspectID": "L1", "first_enrolled": "2026-09-01 05:00:00", "history": [{"x": 1}]},
             "L2": {"ProspectID": "L2", "first_enrolled": None, "history": []}}
    fake = _FakeLSQ(enrolled=leads)
    fetch.enrolled("2026-10-10", str(tmp_path / "e.json"), client=fake, cache=cache)
    rows = {r["ProspectID"]: r for r in json.load(open(tmp_path / "e.json"))}
    assert fake.histories == ["L2", "L3"] and rows["L1"]["history"] == [{"x": 1}]


def test_latest_snapshot_picks_the_newest_pull(tmp_path):
    from analytics.daily_plan import fetch

    for day, name, t in (("2026-10-10", "snap.json", "2026-10-10T17:00:00+00:00"), ("2026-10-11", "snap.json", "2026-10-11T03:30:00+00:00"),
                         ("2026-10-11", "status_1345.json", "2026-10-11T08:15:00+00:00")):
        (tmp_path / day).mkdir(exist_ok=True)
        json.dump({"group": "T", "start": day, "end": day, "fetched_at": t, "users": []}, open(tmp_path / day / name, "w"))
    json.dump([], open(tmp_path / "2026-10-11" / "status_1345_enrolled.json", "w"))
    assert fetch.latest_snapshot(str(tmp_path)).endswith("status_1345.json")
    assert fetch.latest_snapshot(str(tmp_path), start="2026-10-10").endswith("2026-10-10/snap.json")   # the status pull starts too late


def test_lead_flow_types_speed_attempts_and_status_todo():
    from analytics.daily_plan import leadflow

    leads = [dict(l, CreatedOn="2026-10-10 04:30:00", mx_Report_Type="High Priority") if l["ProspectID"] == "L1" else
             dict(l, mx_Assigned_On="2026-10-10 05:00:00") if l["ProspectID"] == "L2" else
             dict(l, CreatedOn="2026-10-10 06:00:00", Source="Inbound Phone call") if l["ProspectID"] == "L3" else l for l in LEADS]
    cs = [call("u1", "L1", "2026-10-10 04:40:00"), call("u1", "L1", "2026-10-10 04:55:00", "Answered", 200),
          call("u0", "L3", "2026-10-10 06:00:00", "Missed", 0, "inbound")]
    s = snap(cs, leads=leads)
    F = {x["lead_id"]: x for x in leadflow.rows(s, "2026-10-10", "2026-10-10")}
    assert set(F) == {"L1", "L2", "L3"}
    assert (F["L1"]["type"], F["L1"]["mins"], F["L1"]["attempts"], F["L1"]["real"]) == ("High Priority", 10, 2, True)
    assert not F["L2"]["new"] and F["L2"]["mins"] is None and F["L2"]["type"] == "Other"
    assert F["L3"]["type"] == "Incoming Missed" and F["L3"]["missed_in"] == 1
    allrow = leadflow.summary(list(F.values()))[-1]
    assert (allrow["group"], allrow["leads"], allrow["dialled"], allrow["never"], allrow["reached"]) == ("All", 3, 1, 2, 1)
    S = status.build(s, "2026-10-10", {}, [], "Tara L")
    assert any("not dialled yet" in t and "Lead2" in t and "Lead3 (Incoming Missed" in t for t in S["todo"])
    assert "New leads today: 3" in status.text_summary(S, "Team")
    assert any("sit with callers who have made no calls today" in t and "Lead2" in t for t in S["todo"])  # Ravi S made no calls
    assert "Incoming Missed" in leadflow.section("t", list(F.values()))
    from integrations.timeutil import IST
    night = datetime(2026, 10, 10, 1, 34, tzinfo=IST)
    assert leadflow.clock_start(night).strftime("%H:%M") == "09:30"
    assert leadflow.clock_start(night.replace(hour=21)).strftime("%d %H:%M") == "11 09:30"


def test_enrolled_cache_reads_every_pull_of_recent_days(tmp_path):
    import os
    import time
    from analytics.daily_plan import fetch

    for day, name, ids in (("2026-10-08", "enrolled.json", ["OLD"]), ("2026-10-10", "enrolled.json", ["A", "B"]),
                           ("2026-10-11", "enrolled.json", ["C"]), ("2026-10-11", "status_1345_enrolled.json", ["D"]),
                           ("2026-10-11", "status_1646_enrolled.json", ["E"]), ("2026-10-11", "status_1822_enrolled.json", ["F"])):
        (tmp_path / day).mkdir(exist_ok=True)
        json.dump([{"ProspectID": i, "first_enrolled": "x"} for i in ids], open(tmp_path / day / name, "w"))
        time.sleep(0.01)
    (tmp_path / "month").mkdir()
    got = fetch.cached_enrolled(str(tmp_path), days=2)
    assert sorted(got) == ["A", "B", "C", "D", "E", "F"]          # all of the last two days, nothing older
    assert os.path.isdir(tmp_path / "month")


def test_a_lead_asking_for_another_day_is_not_due_today():
    sheet = {"L1": {"owner": "Asha K", "tier": "A", "group": 5}, "L3": {"owner": "Asha K", "tier": "A", "group": 1}}
    T = day_stats(snap([]), "2026-10-10", sheet)["team"]
    assert T["a_total"] == 1 and T["a_not_tried"] == ["L3"]


def test_shareable_copy_has_no_phone_numbers_or_caller_lists():
    from analytics.daily_plan import share

    html = ("<html><body><h1>Plan</h1><td>+91-9876543210</td><td>9123456789</td><td>+91-900004053400</td><td>2026-10-12 18:30</td>"
            "<h2>9. Team overview</h2><td>+91-9000000001</td></body></html>")
    out = share.shareable(html)
    assert "9876543210" not in out and "9123456789" not in out and "9000000001" not in out
    assert "900004053400" not in out and "2026-10-12 18:30" in out and "Team overview" not in out
    assert share.lean("<style>x</style><td class='bad' style='a:b'>Lead1</td>\n <td>2</td>") == "<td>Lead1</td><td>2</td>"
    assert share.lean("<td>A <span class=small>(reassigned)</span></td><td>B</td>") == "<td>A <span>(reassigned)</span></td><td>B</td>"


def test_mail_relay_sends_the_real_files(tmp_path, monkeypatch):
    import base64
    monkeypatch.setenv("MAIL_RELAY_URL", "https://script.google.com/macros/s/abc/exec")
    monkeypatch.setenv("MAIL_RELAY_KEY", "k")
    f = tmp_path / "plan.pdf"; f.write_bytes(b"%PDF-1.4 full report")
    calls = []

    class R:
        status_code, text = 200, "ok 5"

    def post(url, json, timeout):
        calls.append((url, json))
        return R()

    assert mailer.send("Plan", "<p>h</p>", "t", [], [str(f)], post=post)
    url, body = calls[0]
    assert body["key"] == "k" and body["to"] == [] and body["attachments"][0]["filename"] == "plan.pdf"
    assert base64.b64decode(body["attachments"][0]["content"]) == b"%PDF-1.4 full report"
    R.text = "forbidden"
    assert not mailer.send("Plan", "<p>h</p>", "t", [], [str(f)], post=post)


def test_restaged_students_and_old_payments_are_not_the_days_enrollments():
    from analytics.daily_plan import review
    reads = {"L1": {"status": "already_student"}, "L3": {"status": "paid_new", "paid_on": "2026-09"}, "L5": {"status": "paid_new", "paid_on": "2026-10-09"}}
    enrolled = [{"ProspectID": lid, "FirstName": lid, "OwnerIdName": "Asha K", "first_enrolled": "2026-10-09 06:00:00", "history": []}
                for lid in ("L1", "L3", "L5", "L7")]
    got = [e["lead_id"] for e in review.enrollments(snap([]), "2026-10-09", enrolled, reads)]
    assert got == ["L5", "L7"]


def test_review_can_drop_record_fixes_from_the_count():
    facts_enr = [{"lead_id": "L1", "name": "Fix", "owner": "Asha K"}, {"lead_id": "L3", "name": "New", "owner": "Asha K"}]
    N = {"not_enrollments": ["L1"]}
    assert [e["name"] for e in facts_enr if e["lead_id"] not in set(N.get("not_enrollments", []))] == ["New"]


def test_revival_pool_history_shortlist_and_ranking(tmp_path):
    from analytics.daily_plan import revival

    leads = [dict(l, CreatedOn="2026-08-05 05:00:00") for l in LEADS]
    leads[1] = dict(leads[1], ProspectStage="Course Enrolled")
    s = snap([], leads=leads)
    ids = revival.pool(s, "2026-07-01", "2026-09-20")
    assert "L2" not in ids and "L1" in ids

    class C:
        def get_lead_activities(self, lid, row_count=200):
            note = "Status{=}Answered{next}Duration{=}300{next}UserId{=}u1"
            acts = [{"Id": f"a{lid}", "EventCode": 22, "CreatedOn": "2026-08-06 05:00:00", "RelatedProspectId": lid,
                     "ActivityFields": {"ActivityEvent_Note": note, "Owner": "u1"}},
                    {"Id": f"e{lid}", "EventCode": 2001, "CreatedOn": "2026-08-06 05:00:00", "RelatedProspectId": lid}]
            return {"ProspectActivities": acts if lid == "L1" else []}

    hist = revival.history(["L1", "L3"], str(tmp_path / "h.jsonl"), workers=2, client=C())
    assert len(hist["L1"]) == 1 and hist["L3"] == []
    again = revival.history(["L1", "L3"], str(tmp_path / "h.jsonl"), client=None)  # resumes from the cache, no client needed
    assert again.keys() == hist.keys()
    rs = Snap(revival.snapshot({"users": USERS, "fetched_at": "2026-10-13T02:00:00+00:00", "leads": leads}, hist), leader="Tara L")
    assert revival.shortlist(rs, ["L1", "L3"]) == ["L1"]
    picked, skipped = revival.focus(rs, ["L1", "L3"], "2026-09-21")
    assert picked == ["L1"] and skipped["no real conversation"] == 1
    assert revival.focus(rs, ["L1"], "2026-08-01")[0] == []          # talked since then: already in the daily reports
    rows = [{"lead_id": "a", "revive": True, "revival_chance": 12, "last_real_talk_date": "2026-08-01"},
            {"lead_id": "b", "revive": True, "revival_chance": 25},
            {"lead_id": "c", "revive": True, "revival_chance": 12, "last_real_talk_date": "2026-09-01"},
            {"lead_id": "d", "revive": False, "revival_chance": 0, "reject_reason": "not interested"}]
    assert [r["lead_id"] for r in revival.ranked(rows)] == ["b", "c", "a"]
    assert revival.band(25) == "High" and revival.band(4) == "Reject"
    assert "L" not in json.dumps(revival.to_state("2026-10-13", [{**r, "lead_id": "L1"} for r in rows[:1]])["leads"])
    assert len(revival.batches([{"owner": "A", "lead_id": str(i), "size": 1000} for i in range(65)])) == 3


def test_dossier_reads_a_zero_second_callfailure_note_as_not_picked_up(tmp_path):
    from analytics.daily_plan import dossiers

    cs = [dict(call("u1", "L1", "2026-10-09 05:00:00"), call_notes="CallFailure"),
          dict(call("u1", "L1", "2026-10-09 06:00:00", "CallFailure"), call_notes="Network error")]
    dossiers.build(snap(cs), "2026-10-09", {"L1"}, {}, {}, str(tmp_path))
    text = (tmp_path / "L1.txt").read_text()
    assert "NotAnswered 0s" in text and "note: CallFailure" not in text   # the lead didn't pick up; not a line fault
    assert "note: Network error" in text                                    # a real failure keeps its note


def test_lead_flow_section_splits_fresh_from_reassigned_leads():
    from analytics.daily_plan import leadflow

    rs = [{"new": True}, {"new": False}, {"new": False}]
    rs = [dict(r, type="Other", owner="Asha K", mins=3, waiting=None, attempts=1, reached=False, real=False, stage="", name="x",
               phone="", arrived=datetime(2026, 10, 10, 10, 0), last="") for r in rs]
    html = leadflow.section("Lead flow", rs, leads=False)
    assert "3 leads came to the team: 1 created in this period (fresh leads) and 2 older leads" in html


def _weekly_fixture(tmp_path):
    from analytics.daily_plan import weekly
    root = tmp_path / "daily"
    for day, calls, fetched in (("2026-10-09", [call("u1", "L1", "2026-10-09 05:00:00", "Answered", 300), call("u1", "L2", "2026-10-09 06:00:00")],
                                 "2026-10-09T16:00:00+00:00"),
                                ("2026-10-10", [call("u1", "L1", "2026-10-09 05:00:00", "Answered", 310),
                                                call("u2", "L2", "2026-10-10 07:00:00", "Answered", 25)], "2026-10-10T16:00:00+00:00")):
        (root / day).mkdir(parents=True)
        json.dump({"group": "Team X", "users": USERS, "leads": LEADS, "calls": calls, "zip_activities": [], "fetched_at": fetched},
                  open(root / day / "snap.json", "w"))
        json.dump([], open(root / day / "snap_enrolled.json", "w"))
    return weekly, Snap(weekly.gather(str(root), "2026-10-09", "2026-10-10"), leader="Tara L")


def test_weekly_review_keeps_each_call_once_and_matches_its_recording(tmp_path):
    weekly, s = _weekly_fixture(tmp_path)
    calls = weekly.connected(s)
    assert [(c["lead_id"], c["duration"]) for c in calls] == [("L1", 310), ("L2", 25)]      # newest pull wins; unanswered left out
    rec = {"start_time": "2026-10-09T05:01:00Z", "call_duration": 305, "transcript": {"text": "caller: fee is 50k. lead: ok"},
           "s3_audio_file_url": "https://acefone/x"}
    tx = {"919000000001": {"sales_call": [rec]}}
    recs = weekly.match(s, calls, tx)
    assert list(recs) == [calls[0]["activity_id"]]
    rows = weekly.ledger(s, calls, recs, set(tx))
    assert [r["transcript"] for r in rows] == ["transcript", "not searched"]
    index = weekly.build_dossiers(s, calls, recs, str(tmp_path / "dossiers"), {})
    text = open(index[0]["path"]).read() if index[0]["lead_id"] == "L1" else open(index[1]["path"]).read()
    assert "fee is 50k" in text and calls[0]["activity_id"] in text


def test_weekly_review_checks_every_call_was_read_and_rolls_up(tmp_path):
    weekly, s = _weekly_fixture(tmp_path)
    calls = weekly.connected(s)
    rows = weekly.ledger(s, calls, {}, set())
    index = [{"lead_id": c["lead_id"], "call_ids": [c["activity_id"]]} for c in calls]
    reads = {"L1": {"lead_id": "L1", "outcome": "open", "primary_reason": "ASK",
                    "calls": [{"call_id": calls[0]["activity_id"], "ask_made": False, "mistake": "no ask"}]}}
    chk = weekly.check(index, reads)
    assert chk["missing_leads"] == ["L2"] and chk["read"] == 1
    reads["L2"] = {"lead_id": "L2", "outcome": "lost", "primary_reason": "L3", "calls": []}
    assert weekly.check(index, reads)["missing_calls"] == [calls[1]["activity_id"]]       # a lead read, one call skipped
    assert weekly.roll_up(rows, reads)["per_caller"]["Asha K"].get("real_read", 0) == 0     # no transcript: the ask isn't judged
    rows[0]["transcript"] = "transcript"
    rows[1]["caller"] = "+91-9000000002"
    R = weekly.roll_up(rows, reads)
    assert R["outcomes"] == {"open": 1, "lost": 1} and dict(R["reasons"]) == {"ASK": 1, "L3": 1}
    assert R["per_caller"]["Asha K"]["real_read"] == 1 and R["mistakes"] == [("no ask", 1)]
    assert "Not a LeadSquared user" in R["per_caller"] and not any("9000" in k for k in R["per_caller"])


def test_weekly_review_counts_an_earlier_enrollment_apart():
    from analytics.daily_plan import weekly
    reads = {"A": {"outcome": "enrolled"}, "B": {"outcome": "enrolled"}, "C": {"outcome": "enrolled"}}
    weekly.settle(reads, {"A": "2026-10-06T10:00:00", "B": "2026-10-09T10:00:00"}, "2026-10-08")
    assert [reads[k]["outcome"] for k in "ABC"] == ["enrolled earlier", "enrolled", "enrolled"]   # C said on a call they paid


def test_weekly_review_gives_a_recording_to_the_call_of_its_length(tmp_path):
    from datetime import timedelta
    weekly, s = _weekly_fixture(tmp_path)
    long_call = weekly.connected(s)[0]                                                      # L1, 310 s at 05:00 UTC
    short = dict(long_call, activity_id="short", duration=5, t=long_call["t"] - timedelta(minutes=1))
    rec = {"start_time": "2026-10-09T04:59:30Z", "call_duration": 305, "transcript": {"text": "caller: fee is 50k"},
           "s3_audio_file_url": "https://acefone/x"}
    recs = weekly.match(s, [short, long_call], {"919000000001": {"sales_call": [rec]}})
    assert list(recs) == [long_call["activity_id"]]                    # closer to the 5 s call, but the length says the long one


def test_weekly_review_spreads_examples_across_callers():
    from analytics.daily_plan import weekly
    leads = [{"owner": "A", "name": f"a{i}"} for i in range(5)] + [{"owner": "B", "name": "unknown"}, {"owner": "B", "name": "b1"}]
    assert [l["name"] for l in weekly._spread(leads, 4)] == ["a0", "b1", "a1", "unknown"]


def test_leader_note_picks_candidates_counts_near_misses_and_renders(tmp_path):
    from analytics.daily_plan import leader_week
    weekly, s = _weekly_fixture(tmp_path)
    reads = {"L1": {"lead_id": "L1", "outcome": "open", "primary_reason": "L6", "owner": "Asha K"},
             "L2": {"lead_id": "L2", "outcome": "lost", "primary_reason": "L3"},
             "L3": {"lead_id": "L3", "outcome": "paying"}, "L4": {"lead_id": "L4", "outcome": "enrolled", "owner": "Asha K"}}
    assert [r["lead_id"] for r in leader_week.candidates(reads)] == ["L3", "L1"]     # unreached and enrolled leads are not near misses
    near = {"L1": {"lead_id": "L1", "name": "Ravi", "owner": "Asha K", "signal": "strong", "saveable": "likely", "tl_lever": "callback_check",
                   "story": "Said yes on Tue; the 2 pm call was never made.", "tl_action": "Check due callbacks at 14:00."},
            "L3": {"lead_id": "L3", "signal": "none", "saveable": "unlikely", "tl_lever": "none"}}
    C = leader_week.near_counts(near)
    assert (C["likely"], C["possible"], C["levers"]) == (1, 0, [("callback_check", 1)])
    F = leader_week.ops(s, "2026-10-09", "2026-10-10", reads)
    assert F["enrolled"] == 1 and set(F["callers"]) == set(s.callers)
    doc = leader_week.html_doc(F, near, {"headline": "x", "moves": [{"title": "Keep callbacks", "do": "d", "why": "w", "examples": ["L1"]}]},
                               {"team": "Team X", "leader": "Tara L", "period": "Fri 9 Oct to Sat 10 Oct", "built": "Sat"})
    assert "Keep callbacks" in doc and doc.count("the 2 pm call was never made") == 1      # an example in a move isn't repeated below


def test_leader_note_separates_bulk_handouts_from_best_leads(monkeypatch):
    from analytics.daily_plan import leader_week
    monkeypatch.setattr(leader_week, "BULK", 3)
    leads = [dict(l, mx_Assigned_On="2026-10-07 04:40:00", mx_Assigned_By="Amit R") for l in LEADS[:4]]          # 4 old leads, one hour
    leads += [dict(LEADS[4], CreatedOn="2026-10-09 05:00:00", mx_Report_Type="High Priority")]                     # a fresh best lead
    s = snap([call("u1", "L1", "2026-10-09 06:00:00", "Answered", 300), call("u1", "L5", "2026-10-09 05:02:00", "Answered", 200)], leads=leads)
    S = leader_week.sources(s, "2026-10-01", "2026-10-10", {"L5": {"outcome": "enrolled"}}, "2026-10-08")
    types = dict(S["types"])
    assert types["Bulk handout of old leads"]["leads"] == 4 and types["High Priority (fresh)"]["enrolled"] == 1
    assert S["handouts"][0]["by"] == "Amit R" and S["handouts"][0]["leads"] == 4
    assert S["best_by_caller"]["Asha K"]["in5"] == 1 and S["bulk_dial_share"] == {}                                 # under 100 dials: no share


def test_missed_calls_sheet_separates_dropped_lines_call_backs_and_waiting_leads(tmp_path):
    from analytics.daily_plan import missed_calls

    def inb(lead, utc, status="Missed", line="LawSikho_Incoming"):
        return dict(call("u1", lead, utc, status, 30 if status == "Answered" else 0, "inbound"), display_number=line, call_notes="CallFailure")
    s = snap([inb("L1", "2026-10-09 05:00:00", "Answered"), inb("L1", "2026-10-09 05:02:00"),       # dropped line, rang straight back
              call("u1", "L1", "2026-10-09 05:06:00", "Answered", 200),
              inb("L3", "2026-10-09 06:00:00"), call("u2", "L3", "2026-10-09 06:20:00", "Answered", 90),  # returned 20 min later
              inb("L5", "2026-10-09 07:00:00", line="Sales_Outbound_Airtel"),                         # rang the outbound number, never returned
              inb("L7", "2026-10-08 07:00:00")])                                                       # another day
    got = {r["lead_id"]: r for r in missed_calls.rows(s, ["2026-10-09"])}
    assert set(got) == {"L1", "L3", "L5"}
    assert got["L1"]["dropped"] and got["L1"]["result"] == "Dropped call, reached again"
    assert got["L3"]["result"] == "Called back and reached" and got["L3"]["back_mins"] == 20 and got["L3"]["back_by"] == "Ravi S"
    assert got["L5"]["result"] == "Not called back" and got["L5"]["line"] == "Rang back the outbound Airtel number"
    assert "90000000" not in json.dumps(missed_calls.by_line(s, ["2026-10-09"]))
    assert missed_calls.by_line(s, ["2026-10-09"])[0] == ("LawSikho incoming line", 3, 1, 2)

    from openpyxl import Workbook
    wb = Workbook()
    workbook.missed_sheet(wb.active, {"label": "Fri 9 Oct", "rows": list(got.values()), "lines": missed_calls.by_line(s, ["2026-10-09"])})
    cells = [c.value for row in wb.active.iter_rows() for c in row if c.value]
    assert "3 missed calls from 3 leads; 1 not called back" in cells and "Not called back" in cells
