from __future__ import annotations

import json
import os

import openpyxl

from analytics import funnel_icp as fi

MAP = {"icps": ["Retired", "Teachers / educators", "General"],
       "pages": {"ret-a": "Retired", "edu-a": "Teachers / educators", "gen-a": "General"},
       "general_segments": {"gen-a": "Brand search"}}


def lead(phone, slug, at, email=""):
    return {"whatsapp": phone, "email": email, "pageUrl": f"https://x/f/{slug}?utm=1", "funnel": slug,
            "captured_ist": at}


def sale(row, phone, email="", paid=0.0):
    return {"row": row, "name": "x", "status": "Disbursed", "phone": phone, "email": email, "paid": paid,
            "paid_raw": paid}


LEADS = [lead("+91 98765 43210", "edu-a", "2026-10-01T10:00"),
         lead("9876543210", "ret-a", "2026-10-02T10:00"),
         lead("9000000001", "ret-a", "2026-10-03T10:00", "A@B.in"),
         lead("9000000003", "gen-a", "2026-10-03T12:00")]
SALES = [sale(2, "9876543210", paid=100.0), sale(3, "1111111111", "a@b.in", 50.0), sale(4, "2222222222", paid=10.0)]
LOGS = [{"call_id": 1, "to": "919876543210", "status": "Contacted", "disposition": "Hangup NEW"},
        {"call_id": 2, "to": "9876543210", "status": "Contacted", "disposition": "Full Conversion NEW"},
        {"call_id": 2, "to": "9876543210", "status": "Contacted", "disposition": "Full Conversion NEW"},
        {"call_id": 3, "to": "9000000001", "status": "No Contact", "disposition": "USER_BUSY"},
        {"call_id": 4, "to": "9000000003", "status": "Contacted", "disposition": "Not Interested NEW"}]


def test_first_funnel_wins_sales_match_by_phone_then_email_and_timepay_best_outcome():
    r = fi.build(LEADS, SALES, MAP, fi.timepay_by_phone(LOGS))
    assert r["people"] == 3
    t = r["icp"]["Teachers / educators"]
    assert (t["leads"], t["dialled"], t["connected"], t["positive"], t["full"], t["sales"], t["paid"]) == \
        (1, 1, 1, 1, 1, 1, 100.0)
    ret = r["icp"]["Retired"]
    assert (ret["dialled"], ret["connected"], ret["positive"], ret["sales"]) == (1, 0, 0, 1)
    assert r["segments"]["Brand search"]["connected"] == 1 and r["segments"]["Brand search"]["positive"] == 0
    assert r["matched"] == {"phone": 1, "email": 1, "not in GrowthX": 1}
    assert all(c[0] != "FAIL" for c in fi.checks(r, MAP, LEADS))


def test_checks_fail_on_unmapped_page_and_missing_segment():
    m = {**MAP, "general_segments": {}}
    r = fi.build(LEADS + [lead("9000000002", "new-page", "2026-10-03T11:00")], SALES, m)
    failed = {c[1] for c in fi.checks(r, m, LEADS) if c[0] == "FAIL"}
    assert "Every lead's landing page is mapped" in failed
    assert "Every General page has a segment, and only General pages" in failed


def test_checks_fail_when_timepay_returns_fewer_calls_than_counted():
    r = fi.build(LEADS, SALES, MAP, {})
    assert ("FAIL", "TimePay: every counted call was read", "read 9 of 10") in fi.checks(r, MAP, LEADS, 10, 9)


def test_read_sales_keeps_sale_statuses_and_float_phones(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bootcamp | X"
    ws.append(["Bootcamp Lead", "Date", "Student name", "Phone No.", "Email Id", "Money received ",
               "Status of the lead"])
    ws.append(["", "", "A", 919876543210.0, "A@x.in", 5000, "Disbursed"])
    ws.append(["", "", "B", 919000000000.0, "", 0, "Not interested"])
    path = os.path.join(tmp_path, "s.xlsx")
    wb.save(path)
    got = fi.read_sales(path, "Bootcamp | X")
    assert [(s["row"], s["phone"], s["email"], s["paid"]) for s in got] == [(2, "9876543210", "a@x.in", 5000.0)]


def test_written_report_rereads_to_the_same_totals(tmp_path):
    r = fi.build(LEADS, SALES, MAP, fi.timepay_by_phone(LOGS))
    path = os.path.join(tmp_path, "r.xlsx")
    fi.write_xlsx(r, fi.checks(r, MAP, LEADS), path, "t", "fresh")
    assert fi.verify_written(path, r)


def test_shipped_maps_are_consistent():
    for name in os.listdir(fi.MAPS_DIR):
        m = json.load(open(os.path.join(fi.MAPS_DIR, name), encoding="utf-8"))
        assert set(m["pages"].values()) <= set(m["icps"]), name
        assert set(m["general_segments"]) == {s for s, v in m["pages"].items() if v == "General"}, name


def test_caller_list_priorities_tabs_and_icp_choice(tmp_path):
    from analytics import funnel_icp_callers as fc

    m = {**MAP, "playbook": {"Retired": {"open_with": "R line", "story": "R story", "day": "Day 1"}},
         "stated_icp_codes": {"E": "Retired", "F": "General"}}
    leads = LEADS + [lead("9000000004", "edu-a", "2026-10-03T13:00"), lead("9000000005", "ret-a", "2026-10-03T14:00")]
    logs = LOGS + [{"call_id": 5, "to": "9000000004", "status": "Contacted", "disposition": "Not Interested NEW"},
                   {"call_id": 6, "to": "9000000005", "status": "Contacted", "disposition": "Call Back NEW",
                    "summary": "Asked to call after 6 PM", "start_time": "2026-10-04T10:00:00"}]
    tp = fi.timepay_by_phone(logs)
    r = fi.build(leads, SALES, m, tp)
    stated = {"9000000003": "Retired", "9000000004": "General"}
    rows = {x["Phone"]: x for x in fc.caller_rows(leads, r, m, tp, stated)}
    assert rows["9876543210"]["tab"] == "Already enrolled" and rows["9876543210"]["Priority"] == ""
    no = rows["9000000004"]  # said no: stays on the call list, last
    assert (no["tab"], no["prio"], no["Said yes to the AI agent?"]) == ("Call list", 6, "No")
    assert no["ICP"] == "Teachers / educators"  # call "General" never overrides the ad
    assert rows["9000000003"]["ICP"] == "Retired" and rows["9000000003"]["Check"]
    assert rows["9000000003"]["Open with"] == "R line"
    assert rows["9000000001"]["Said yes to the AI agent?"] == "Not reached" and rows["9000000001"]["prio"] == 5
    cb = rows["9000000005"]  # callback; Retired sold 1 of 2 leads, above the campaign average
    assert (cb["prio"], cb["Said yes to the AI agent?"], cb["AI call summary"], cb["Last AI call (IST)"]) == \
        (3, "Callback", "Asked to call after 6 PM", "2026-10-04T10:00:00")
    assert cb["ICP sales so far"] == "1 of 2 leads bought (50.0%)"
    assert all(c[0] == "PASS" for c in fc.caller_checks(list(rows.values()), r))
    path = os.path.join(tmp_path, "c.xlsx")
    fc.write_caller_xlsx(list(rows.values()), path, "t", "f", m, r)
    assert openpyxl.load_workbook(path).sheetnames == ["How to use", "Call list", "Already enrolled"]


def test_shipped_maps_have_a_playbook_for_every_icp():
    for name in os.listdir(fi.MAPS_DIR):
        m = json.load(open(os.path.join(fi.MAPS_DIR, name), encoding="utf-8"))
        assert set(m["playbook"]) == set(m["icps"]), name
        assert set(m["stated_icp_codes"].values()) <= set(m["icps"]), name


def test_every_tab_is_written_as_csv(tmp_path):
    r = fi.build(LEADS, SALES, MAP, fi.timepay_by_phone(LOGS))
    path = os.path.join(tmp_path, "r.xlsx")
    fi.write_xlsx(r, fi.checks(r, MAP, LEADS), path, "t", "fresh")
    names = sorted(os.path.basename(p) for p in fi.xlsx_to_csvs(path))
    assert names == ["r_By_ICP.csv", "r_By_funnel.csv", "r_Checks.csv", "r_Spot_check.csv"]
    text = open(os.path.join(tmp_path, "r_By_ICP.csv"), encoding="utf-8-sig").read()
    assert text.startswith("t") and "Teachers / educators" in text
