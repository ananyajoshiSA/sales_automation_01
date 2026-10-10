"""Audit of every open collection lead: its chance of being collected, what blocks it, and the pipeline it adds up to.

The audit takes every lead in the collection pool (``analytics.bootcamp_collections``) booked on or after
``since`` that is still open or parked at "May buy later", and for each one:

* **Where it stands now:** stage, days since booking, days since anyone last spoke to the lead (a real
  conversation, 120 s+), its missed calls not returned, follow-up overdue, untouched since hand-over, and
  notes that say the lead paid or the loan went through while the stage still reads open.
* **Its chance, from history:** leads booked at least 45 days before the data was pulled are looked at as
  they stood 2, 5, 10, 17, 25 and 35 days after booking. Among those still open, the share collected in the
  next 3, 14, 21 and 45 days is the chance for an open lead today in the same position (bootcamp or community,
  age, stage, how recently someone spoke to the lead). A position with fewer than 30 past examples falls
  back to a coarser one (no contact recency, then no stage); the level used is shown. Each chance has a 90%
  range (Wilson). These are measured shares of past leads, not promises.
* **A file to read:** a short history since booking (stage changes with their comments, answered calls,
  unanswered dials per day, missed calls from the lead, owner changes, form notes, Zipteams summaries),
  handed out in batches under ``<out_dir>/reading/<data time>/`` with INSTRUCTIONS.md. Claude reads each lead in a
  Claude Code session (no model API) and writes its situation, blocker, outlook, next step, what to say and
  handling issues to ``read_NN.jsonl``; ``check`` validates a file, and the next run merges every reading
  that passes into the call list. Given the transcripts of the open leads
  (``scripts/fetch_collection_transcripts.py``), the round is a second reading that also sets each lead's
  objection, where it is stuck and how likely it is to pay by month end.

The pipeline adds the chances per caller, team and kind; its 90% range comes from 2,000 simulated runs that
draw each lead's chance from its range. With the leads already collected it gives the projected collection
rate of the pool booked since ``since`` against the 75% and 80% targets.

Writes ``open_leads.csv``, ``by_kind_team.csv``, ``by_kind_caller.csv``, ``report.json`` and
``collection_audit.xlsx`` (a summary sheet and a call list per team, in calling order) to ``out_dir``.

    python -m analytics.collection_audit data/coll_now exports/collection_audit 2026-09-01 [transcripts.jsonl]
    python -m analytics.collection_audit check <round folder>/read_01.jsonl <round folder>/batch_01.jsonl
"""

from __future__ import annotations

import csv
import json
import math
import os
import random
import re
import sys
from collections import Counter, defaultdict
from collections.abc import Mapping
from datetime import date, datetime, timedelta

from analytics.bootcamp_collections import (BOOKED, COLLECTED, LOST, _call, _data, _fields,
                                            booking_start, build, is_community, webinar_weekend)
from analytics.definitions import REAL_CONVERSATION_SECS
from analytics.lead_priority import strip_html
from integrations.timeutil import IST, utc

SNAPSHOT_DAYS = (2, 5, 10, 17, 25, 35)
HORIZONS = (3, 14, 21, 45)   # 21: from a 10 Oct pull to the end of the month
MIN_CELL = 30
Z90 = 1.645
TARGETS = (75, 80)

STAGE_GROUP = {
    "booking fees received": "Untouched (still Booking fees received)",
    "call not picking up": "Not picking up", "call not connected": "Not picking up",
    "call back later": "Call back later",
    "follow up for closure": "Follow-up for closure",
    "loan pending": "Loan pending",
    "may buy later": "May buy later",
}
PAID_WORDS = re.compile(r"\bpaid\b|payment (is )?(done|received|completed|made)|loan (process )?(is )?(done|approved|disbursed"
                        r"|sanctioned|completed)|disburs|\butr\b|transaction id", re.I)
SYSTEM_COMMENT = re.compile(r'^\s*\{"ActionType"')


def stage_group(stage: str) -> str:
    return STAGE_GROUP.get((stage or "").strip().lower(), "Other (counselled, roadmap, discovery)")


def age_bucket(days: float) -> str:
    for limit, label in ((3, "0-3 days"), (7, "4-7 days"), (14, "8-14 days"), (21, "15-21 days"), (30, "22-30 days")):
        if days <= limit:
            return label
    return "31+ days"


def contact_bucket(days: float | None) -> str:
    if days is None:
        return "never spoke since booking"
    return "spoke in last 2 days" if days <= 2 else "spoke 3-7 days ago" if days <= 7 else "last spoke 8+ days ago"


def position(acts: list[dict], booked: datetime, t: datetime) -> dict:
    """Where a lead stood at ``t``: its stage then and days since the last real conversation."""
    stage = BOOKED
    for a in acts:
        at = utc(a.get("CreatedOn"))
        if a.get("EventCode") == 3002 and at and booked <= at <= t:
            stage = (_data(a).get("CurrentStage") or "").strip() or stage
    spoke = [c["t"] for c in (_call(a) for a in acts if a.get("EventCode") in (21, 22))
             if c["t"] and booked <= c["t"] <= t and c["answered"] and c["dur"] >= REAL_CONVERSATION_SECS]
    return {"stage": stage, "days_open": (t - booked).total_seconds() / 86400,
            "days_since_spoke": (t - max(spoke)).total_seconds() / 86400 if spoke else None}


def cell_keys(kind: str, pos: dict) -> list[tuple]:
    """Finest to coarsest position keys."""
    age, stage = age_bucket(pos["days_open"]), stage_group(pos["stage"])
    return [(kind, age, stage, contact_bucket(pos["days_since_spoke"])), (kind, age, stage), (kind, age)]


def history_table(leads: list[dict], hist: Mapping[str, list], now: datetime) -> dict[tuple, list[int]]:
    """For every position key: [snapshots, collected within each of HORIZONS days]."""
    table: dict[tuple, list[int]] = defaultdict(lambda: [0] * (1 + len(HORIZONS)))
    for lead in leads:
        acts = hist.get(lead["ProspectID"])
        if not acts:
            continue
        acts = sorted((a for a in acts if utc(a.get("CreatedOn"))), key=lambda a: a["CreatedOn"])
        stages = [(utc(a["CreatedOn"]), _data(a)) for a in acts if a.get("EventCode") == 3002]
        booked = booking_start([t for t, d in stages if (d.get("CurrentStage") or "").strip().lower() == BOOKED])
        if not booked:
            continue
        kind = "Community" if is_community(lead.get("mx_Bootcamp_collections")) else "Bootcamp"
        collected = next((t for t, d in stages if t >= booked and d.get("CurrentStage") in COLLECTED), None)
        for day in SNAPSHOT_DAYS:
            t = booked + timedelta(days=day)
            if t + timedelta(days=max(HORIZONS)) > now or (collected and collected <= t):
                continue
            pos = position(acts, booked, t)
            if pos["stage"] in LOST:
                continue
            hits = [bool(collected and collected <= t + timedelta(days=h)) for h in HORIZONS]
            for key in cell_keys(kind, pos):
                row = table[key]
                row[0] += 1
                for i, hit in enumerate(hits, 1):
                    row[i] += hit
    return dict(table)


def wilson(k: int, n: int) -> tuple[float, float]:
    if not n:
        return 0.0, 1.0
    p = k / n
    d = 1 + Z90 ** 2 / n
    c = (p + Z90 ** 2 / (2 * n)) / d
    h = Z90 * math.sqrt(p * (1 - p) / n + Z90 ** 2 / (4 * n * n)) / d
    return max(0.0, c - h), min(1.0, c + h)


def chance(table: Mapping[tuple, list[int]], kind: str, pos: dict) -> dict:
    """The lead's chance of being collected within each horizon, from the finest position with enough history."""
    keys = cell_keys(kind, pos)
    level = next((i for i, k in enumerate(keys) if table.get(k, [0])[0] >= MIN_CELL), len(keys) - 1)
    n, *hits = table.get(keys[level], [0] * (1 + len(HORIZONS)))
    out = {}
    for h, k in zip(HORIZONS, hits):
        lo, hi = wilson(k, n)
        out.update({f"chance_{h}d": round(k / n, 3) if n else None, f"chance_{h}d_low": round(lo, 3),
                    f"chance_{h}d_high": round(hi, 3)})
    return {**out, "chance_from": ("age, stage and contact", "age and stage", "age only")[level], "chance_examples": n}


def simulate(rows: list[dict], key: str, runs: int = 2000, seed: int = 7) -> tuple[float, int, int]:
    """Expected number collected and its 90% range: each run draws each lead's chance from its range."""
    rng = random.Random(seed)
    totals = []
    for _ in range(runs):
        hit = 0
        for r in rows:
            lo, hi = r[f"{key}_low"], r[f"{key}_high"]
            hit += rng.random() < rng.uniform(lo, hi)
        totals.append(hit)
    totals.sort()
    return round(sum(r[key] or 0 for r in rows), 1), totals[int(0.05 * runs)], totals[int(0.95 * runs) - 1]


def flags(view: dict, acts: list[dict], booked: datetime, now: datetime) -> list[str]:
    """Plain-language problems a team leader should see on this lead."""
    out = []
    comments = [(_data(a).get("Comment") or "").strip() for a in acts if a.get("EventCode") == 3002
                and utc(a.get("CreatedOn")) and utc(a["CreatedOn"]) >= booked]
    notes = comments + [(_fields(a).get("ActivityEvent_Note") or "") for a in acts if a.get("EventCode") == 103
                        and utc(a.get("CreatedOn")) and utc(a["CreatedOn"]) >= booked]
    if any(PAID_WORDS.search(n) for n in notes if n and not SYSTEM_COMMENT.match(n)):
        out.append("Notes say paid or loan done, but the stage is still open: check the payment")
    if view["stage"].strip().lower() == BOOKED and not view["real_convs"]:
        out.append("Untouched since hand-over: no real conversation yet")
    if view["days_since_last_dial"] is None:
        out.append("Never dialled by anyone since booking")
    elif view["days_since_last_dial"] > 3:
        out.append(f"No dial for {view['days_since_last_dial']:.0f} days")
    if view["inbound_unreturned_days"]:
        out.append(f"Lead called and was not called back the same day ({view['inbound_unreturned_days']} day(s))")
    if view["followup_overdue"]:
        out.append("Follow-up date has passed with no dial")
    if view["stage"].strip().lower() == "loan pending":
        loan_since = max((utc(a["CreatedOn"]) for a in acts if a.get("EventCode") == 3002
                          and (_data(a).get("CurrentStage") or "").strip().lower() == "loan pending"), default=None)
        if loan_since and now - loan_since > timedelta(days=7):
            out.append(f"In Loan pending for {(now - loan_since).days} days")
    if view["max_dials_one_day"] >= 5 and not view["real_convs"]:
        out.append("Dialled 5+ times in a day with no real conversation yet")
    return out


def dossier(lead: dict, view: dict, acts: list[dict], booked: datetime, now: datetime) -> str:
    """The lead's history since booking, short enough to read in a minute (IST times)."""
    lines = []
    unanswered: Counter = Counter()

    def flush() -> None:
        for (day, by), n in sorted(unanswered.items()):
            lines.append(f"{day} {n} unanswered dial(s) by {by or 'unknown'}")
        unanswered.clear()

    for a in acts:
        t = utc(a.get("CreatedOn"))
        if not t or t < booked - timedelta(minutes=5):
            continue
        when = t.astimezone(IST).strftime("%d %b %a %H:%M")
        code = a.get("EventCode")
        if code in (21, 22):
            c = _call(a)
            if code == 22 and not c["answered"]:
                unanswered[(t.astimezone(IST).strftime("%d %b"), c["by"])] += 1
                continue
            flush()
            if code == 22:
                lines.append(f"{when} call by {c['by'] or 'unknown'}: answered, {c['dur'] // 60} min {c['dur'] % 60} s")
            else:
                lines.append(f"{when} LEAD CALLED IN: {'answered' if c['answered'] else 'missed'}")
            continue
        flush()
        if code == 3002:
            d = _data(a)
            comment = (d.get("Comment") or "").strip()
            comment = "" if SYSTEM_COMMENT.match(comment) else comment
            lines.append(f"{when} stage {d.get('PreviousStage')} -> {d.get('CurrentStage')}" + (f": {comment[:300]}" if comment else ""))
        elif code == 3001:
            lines.append(f"{when} owner -> {_data(a).get('CurrentOwner')}")
        elif code == 103:
            f = _fields(a)
            note = (f.get("ActivityEvent_Note") or "").strip()
            due = utc(f.get("mx_Custom_1"))
            lines.append(f"{when} form {f.get('Status') or ''}" + (f", callback {due.astimezone(IST).strftime('%d %b %H:%M')}" if due else "")
                         + (f": {note[:200]}" if note else ""))
        elif code == 237:
            f = _fields(a)
            text = strip_html(f.get("ActivityEvent_Note") or "")[:400]
            lines.append(f"{when} Zipteams summary (intent {f.get('mx_Custom_1') or '?'}): {text}")
    flush()
    head = (f"{view['kind']} | {view['course']} | booked {view['booked_ist']} IST ({view['days_since_booking']:.0f} days ago) | "
            f"caller {view['caller']} ({view['team']}) | owner now {view['owner']} | stage now {view['stage']} | "
            f"course fee on lead: {lead.get('mx_Course_Fees') or 'not set'}")
    return head + "\n" + "\n".join(lines[-40:])


def _load(data_dir: str) -> tuple[dict, list[dict], dict[str, list]]:
    meta = json.load(open(os.path.join(data_dir, "meta.json")))
    leads = json.load(open(os.path.join(data_dir, "tagged_leads.json")))
    hist = {}
    for line in open(os.path.join(data_dir, "hist.jsonl")):
        d = json.loads(line)
        if d.get("activities") is not None and not d.get("error"):
            hist[d["lead_id"]] = d["activities"]
    return meta, leads, hist


def audit(leads: list[dict], hist: Mapping[str, list], team_of_owner: Mapping[str, str], callers, now: datetime,
          since: str) -> tuple[list[dict], list[dict], dict, list[dict]]:
    """(open lead rows, every view in the pool since ``since``, history table, every view)."""
    views, _ = build(leads, hist, team_of_owner, now, callers)
    pool = [v for v in views if v["booked_ist"][:10] >= since]
    table = history_table(leads, hist, now)
    by_id = {l["ProspectID"]: l for l in leads}
    rows = []
    for v in pool:
        if v["outcome"] not in ("open", "deferred"):
            continue
        acts = sorted((a for a in hist[v["lead_id"]] if utc(a.get("CreatedOn"))), key=lambda a: a["CreatedOn"])
        pos = position(acts, v["booked"], now)
        lead = by_id[v["lead_id"]]
        rows.append({
            "lead_id": v["lead_id"], "kind": v["kind"], "team": v["team"], "caller": v["caller"], "owner_now": v["owner"],
            "course": v["course"], "booked_ist": v["booked_ist"], "days_open": round(pos["days_open"], 1),
            "stage": v["stage"], "stage_group": stage_group(v["stage"]),
            "days_since_spoke": round(pos["days_since_spoke"], 1) if pos["days_since_spoke"] is not None else None,
            "days_since_last_dial": v["days_since_last_dial"], "dials": v["dials"], "real_convs": v["real_convs"],
            "course_fee": lead.get("mx_Course_Fees") or "", **chance(table, v["kind"], pos),
            "flags": flags(v, acts, v["booked"], now), "dossier": dossier(lead, v, acts, v["booked"], now),
        })
    return rows, pool, table, views


def month_end(now: datetime) -> datetime:
    """23:59 IST on the last day of ``now``'s IST month."""
    d = now.astimezone(IST)
    first_next = (d.replace(day=28) + timedelta(days=4)).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    return first_next - timedelta(minutes=1)


def new_bookings_outlook(views: list[dict], now: datetime, weeks: int = 5) -> list[dict]:
    """Estimate, per kind and team, of collections from bookings still to come before month end.

    Weekly volume: bookings in each of the last ``weeks`` complete webinar weeks (Saturday to Friday). Rate: share
    of leads booked 45 to 150 days ago that were collected within the days a new booking would have left in the
    month (from the Monday of its webinar weekend). Weeks still to come: this week's weekend onwards, if its
    Monday falls before month end; bookings already in from this week are counted off it.
    """
    end = month_end(now)
    this_week = webinar_weekend(now)
    weekends, w = [], date.fromisoformat(this_week)
    while datetime.combine(w + timedelta(days=2), datetime.min.time(), IST) < end:
        weekends.append(w)
        w += timedelta(days=7)
    past = sorted({v["weekend"] for v in views if v["weekend"] < this_week})[-weeks:]
    old = [v for v in views if timedelta(days=45) <= now - v["booked"] <= timedelta(days=150)]
    out = []
    for key in sorted({(v["kind"], v["team"]) for v in views}):
        vol = [sum(1 for v in views if (v["kind"], v["team"]) == key and v["weekend"] == wk) for wk in past]
        if not vol or not max(vol):
            continue
        base = [v for v in old if (v["kind"], v["team"]) == key]
        seen = sum(1 for v in views if (v["kind"], v["team"]) == key and v["weekend"] == this_week)
        exp = lo = hi = 0.0
        detail = []
        for wk in weekends:
            days = (end - datetime.combine(wk + timedelta(days=2), datetime.min.time(), IST)).days
            k = sum(1 for v in base if v["days_to_collect"] is not None and v["days_to_collect"] <= days)
            rate = k / len(base) if base else 0.0
            r_lo, r_hi = wilson(k, len(base))
            left = max(0, sorted(vol)[len(vol) // 2] - (seen if wk.isoformat() == this_week else 0))
            exp += left * rate
            lo += max(0, min(vol) - (seen if wk.isoformat() == this_week else 0)) * r_lo
            hi += max(0, max(vol) - (seen if wk.isoformat() == this_week else 0)) * r_hi
            detail.append(f"{wk:%d %b}: ~{left} bookings x {100 * rate:.0f}% within {days} days")
        out.append({"kind": key[0], "team": key[1], "weekly_bookings": vol, "expected": round(exp, 1),
                    "range": f"{lo:.0f}-{hi:.0f}", "weeks": detail})
    return out


def pipeline(rows: list[dict], pool: list[dict], *keys: str) -> list[dict]:
    """Per group: pool booked since the start date, collected so far, open, expected collections and the projected rate."""
    g_rows: dict[tuple, list] = defaultdict(list)
    g_pool: dict[tuple, list] = defaultdict(list)
    for r in rows:
        g_rows[tuple(r[k] for k in keys)].append(r)
    for v in pool:
        g_pool[tuple(v[k] for k in keys)].append(v)
    out = []
    for k in sorted(g_pool, key=lambda k: tuple(map(str, k))):
        p, r = g_pool[k], g_rows.get(k, [])
        n, done = len(p), sum(1 for v in p if v["outcome"] == "collected")
        row = {**dict(zip(keys, k)), "pool": n, "collected": done, "lost": sum(1 for v in p if v["outcome"] == "lost"),
               "open": len(r), "collected_%_now": round(100 * done / n, 1)}
        for h in HORIZONS:
            e, lo, hi = simulate(r, f"chance_{h}d")
            row.update({f"expected_next_{h}d": e, f"range_{h}d": f"{lo}-{hi}"})
            if h == 45:
                lo45, hi45, e45 = lo, hi, e
        row.update({
               "projected_%_45d": round(100 * (done + e45) / n, 1),
               "projected_range_%_45d": f"{100 * (done + lo45) / n:.0f}-{100 * (done + hi45) / n:.0f}"})
        for t in TARGETS:
            need = max(0, math.ceil(t * n / 100) - done)
            row[f"needed_for_{t}%"] = need
            row[f"needed_for_{t}%_share_of_open"] = f"{100 * need / len(r):.0f}%" if r else "-"
        out.append(row)
    return out


READING_SIZE = 80
BLOCKERS = (
    "May have paid already - check", "Loan or EMI paperwork in progress", "Wants a loan or EMI option",
    "Arranging money / promised a date", "Cannot afford now", "Family or spouse to decide",
    "Busy, travelling or personal reasons", "Doubts about the course or its value", "Wants a refund or to drop",
    "Not reachable", "Not yet properly counselled", "No clear reason recorded",
)
OUTLOOKS = ("Check payment first", "Ready to close", "Workable with follow-up", "Long shot", "Likely lost")
PROCESS_ISSUES = (
    "Slow first call", "Lead's call not returned", "Long gap with no follow-up", "Promised callback missed",
    "Very short calls, no counselling", "Over-dialled without contact", "Loan stalled", "Stage not updated",
    "No dated next step",
)
READING_TEXT = {"situation": 240, "next_action": 220, "what_to_say": 280, "when": 60}
OUTLOOK_ORDER = {o: i for i, o in enumerate(OUTLOOKS)}
STUCK_AT = ("May have paid - check", "Payment link sent, not paid", "Promised to pay on a date", "Loan approval or KYC",
            "Loan or EMI documents", "Choosing how to pay", "Counselled, still deciding", "Not yet counselled",
            "Unreachable", "Wants to defer, refund or drop")
MONTH_END = ("Very likely", "Likely", "Possible", "Unlikely", "Very unlikely")
V2_TEXT = {"objection": 200, "month_end_reason": 220}
V2_SIZE = 25
TRANSCRIPT_CHARS = 10000     # a call longer than this keeps its opening and its last part (payment talk comes late)
LEAD_CHARS = 20000           # whole conversation for most leads (median 14,000 characters since booking, 10 Oct)

INSTRUCTIONS_V2 = """
## Round 2: transcripts, objection, where it is stuck, month-end likelihood

Each lead in this round also carries "transcripts": its recorded calls since around the booking, oldest first,
up to {lead_chars:,} characters (newest calls kept first; API date, which can be 5 h 30 min off; a call over
{call_chars:,} characters is cut in the middle, marked [...]). Read them in full with the dossier. Read them with the dossier:
what the lead actually said outweighs a short stage note. Add four fields to each reading:

- "objection": the lead's own objection or concern in plain words (max {objection} characters), e.g. "Wants
  6-month EMI; card limit too low"; "none raised" if there is none.
- "stuck_at": where the lead is stuck, exactly one of: {stuck}.
- "month_end": how likely the lead is to pay the balance by 31 October, from everything read, exactly one of:
  {month_end}. Judge from the lead's words and behaviour (a payment date before month end, documents done,
  answering and engaged) versus deferrals past October, silence, refund talk.
- "month_end_reason": one sentence on why (max {month_end_reason} characters).

The earlier fields stay as described above; update them where the transcripts change the picture.
"""


def transcript_excerpts(calls: list[dict], booked: datetime) -> list[dict]:
    """Recorded calls since two days before booking, newest kept first up to LEAD_CHARS, returned oldest first;
    a call over TRANSCRIPT_CHARS keeps its opening and its end."""
    keep = []
    for c in calls:
        try:
            at = datetime.fromisoformat(c["start_api"]) if c.get("start_api") else None
        except ValueError:
            at = None
        text = (c.get("transcript") or "").strip()
        if not text or not at or at < booked - timedelta(days=2):
            continue
        if len(text) > TRANSCRIPT_CHARS:
            head = TRANSCRIPT_CHARS // 4
            text = text[:head] + " [...] " + text[-(TRANSCRIPT_CHARS - head):]
        keep.append({"date": at.astimezone(IST).strftime("%d %b"), "agent": c.get("agent") or "",
                     "minutes": round((c.get("duration") or 0) / 60, 1), "text": text, "_t": at})
    keep.sort(key=lambda c: c["_t"], reverse=True)
    out, used = [], 0
    for c in keep:
        if used + len(c["text"]) > LEAD_CHARS:
            if out:
                break
            c["text"] = c["text"][:LEAD_CHARS]
        out.append(c)
        used += len(c["text"])
    out.sort(key=lambda c: c["_t"])
    return [{k: v for k, v in c.items() if k != "_t"} for c in out]

INSTRUCTIONS = """# Reading the open collection leads

You are Claude, setting the next step for each open collection lead of LawSikho's two collection teams
(Elite Changemakers under Mayur Sachdeva, Team Puja Malik under Puja Malik). Each lead paid a booking amount
during a bootcamp (Rs 3,000; the balance is about Rs 55,000-60,000) or one of the community's sales webinars,
and the team has to collect the balance. Loans and EMIs go through finance partners; notes often name the
loan desk (e.g. "docs shared with Dakshita / Sakshi").

Data as of {as_of} IST. Privacy: the files hold real leads. Keep them in this folder; never quote them in chat,
logs or commits.

For each line of your batch file (`batch_NN.jsonl`: "lead_id", "flags" and "dossier"), read the dossier and the
flags and write one JSON object per lead, one per line, to `read_NN.jsonl` in this folder:

- "lead_id": copied exactly.
- "situation": one plain sentence on where this lead stands and why (max {situation} characters). No names,
  phone numbers or emails: say "the lead", "the spouse", "the caller".
- "blocker": the main thing between this lead and payment, exactly one of: {blockers}.
- "outlook": exactly one of: {outlooks}. "Check payment first" when the notes say paid, UTR, loan done or
  disbursed while the stage is open. "Ready to close" when only a last step is left (link sent, KYC left, a
  promised date close by). "Likely lost" when the lead asked for a refund or to drop and nothing since says
  otherwise.
- "next_action": what the caller does next, starting with a verb (max {next_action} characters), e.g.
  "Call today 6-7 PM as the lead asked; confirm the video KYC is done with the loan desk".
- "what_to_say": the angle for that call in the caller's words (max {what_to_say} characters): answer the
  blocker, never pressure. Counsel first where the lead was never properly counselled.
- "when": "today", "tomorrow", a date like "12 Oct" that the lead asked for, or "before calling: check payment"
  (max {when} characters).
- "process_issues": a list (may be empty) of what went wrong in handling this lead, each exactly one of:
  {issues}.

Use only what the dossier shows; when it is thin, say so in "situation" and choose "No clear reason recorded".
When you finish, run from the repository root:

    python -m analytics.collection_audit check {folder}/read_NN.jsonl {folder}/batch_NN.jsonl

and fix every problem it lists until it prints "ok".
"""


def round_version(folder: str) -> int:
    path = os.path.join(folder, "round.json")
    return json.load(open(path)).get("version", 1) if os.path.exists(path) else 1


def check_reading(read_path: str, batch_path: str) -> list[str]:
    """Problems with a reading file: every lead of the batch read once, every label from its list."""
    v2 = round_version(os.path.dirname(os.path.abspath(batch_path))) >= 2
    want = [json.loads(line)["lead_id"] for line in open(batch_path, encoding="utf-8") if line.strip()]
    problems, seen = [], Counter()
    for n, line in enumerate(open(read_path, encoding="utf-8"), 1):
        if not line.strip():
            continue
        try:
            r = json.loads(line)
        except ValueError:
            problems.append(f"line {n}: not JSON")
            continue
        lid = r.get("lead_id")
        seen[lid] += 1
        if lid not in want:
            problems.append(f"line {n}: lead_id {lid} is not in the batch")
        for key, limit in READING_TEXT.items():
            if not isinstance(r.get(key), str) or not r[key].strip():
                problems.append(f"line {n}: {key} missing")
            elif len(r[key]) > limit:
                problems.append(f"line {n}: {key} longer than {limit} characters")
        if r.get("blocker") not in BLOCKERS:
            problems.append(f"line {n}: blocker {r.get('blocker')!r} not in the list")
        if r.get("outlook") not in OUTLOOKS:
            problems.append(f"line {n}: outlook {r.get('outlook')!r} not in the list")
        issues = r.get("process_issues")
        if not isinstance(issues, list) or any(i not in PROCESS_ISSUES for i in issues):
            problems.append(f"line {n}: process_issues must be a list from the list")
        if v2:
            for key, limit in V2_TEXT.items():
                if not isinstance(r.get(key), str) or not r[key].strip():
                    problems.append(f"line {n}: {key} missing")
                elif len(r[key]) > limit:
                    problems.append(f"line {n}: {key} longer than {limit} characters")
            if r.get("stuck_at") not in STUCK_AT:
                problems.append(f"line {n}: stuck_at {r.get('stuck_at')!r} not in the list")
            if r.get("month_end") not in MONTH_END:
                problems.append(f"line {n}: month_end {r.get('month_end')!r} not in the list")
    problems += [f"lead {lid} read {k} times" for lid, k in seen.items() if k > 1]
    problems += [f"lead {lid} not read" for lid in want if not seen[lid]]
    return problems


def write_reading_round(rows: list[dict], reading_dir: str, as_of: str,
                        transcripts: Mapping[str, list] | None = None) -> str:
    """A round folder (named for the data time) with batches of lead files and the instructions. With
    ``transcripts`` (lead_id -> excerpts) it is a round-2 reading, which also sets the objection, where the lead
    is stuck and its month-end likelihood."""
    v2 = transcripts is not None
    folder = os.path.join(reading_dir, as_of.replace("-", "").replace(" ", "-").replace(":", "") + ("-v2" if v2 else ""))
    os.makedirs(folder, exist_ok=True)
    for name in os.listdir(folder):
        if name.startswith("batch_"):
            os.remove(os.path.join(folder, name))
    size = V2_SIZE if v2 else READING_SIZE
    batches = [rows[i:i + size] for i in range(0, len(rows), size)]
    for i, batch in enumerate(batches, 1):
        with open(os.path.join(folder, f"batch_{i:02d}.jsonl"), "w", encoding="utf-8") as fh:
            for r in batch:
                item = {"lead_id": r["lead_id"], "flags": r["flags"], "dossier": r["dossier"]}
                if v2:
                    item["transcripts"] = transcripts.get(r["lead_id"], [])
                fh.write(json.dumps(item, ensure_ascii=False) + "\n")
    text = INSTRUCTIONS.format(as_of=as_of, folder=folder, blockers="; ".join(BLOCKERS), outlooks="; ".join(OUTLOOKS),
                               issues="; ".join(PROCESS_ISSUES), **READING_TEXT)
    if v2:
        text += INSTRUCTIONS_V2.format(lead_chars=LEAD_CHARS, call_chars=TRANSCRIPT_CHARS, stuck="; ".join(STUCK_AT),
                                       month_end="; ".join(MONTH_END), **V2_TEXT)
    json.dump({"version": 2 if v2 else 1, "as_of": as_of}, open(os.path.join(folder, "round.json"), "w"))
    open(os.path.join(folder, "INSTRUCTIONS.md"), "w", encoding="utf-8").write(text)
    return folder


def load_readings(reading_dir: str) -> dict[str, dict]:
    """Readings from every round's read_NN.jsonl that passes its check against its batch; a later round wins.
    Each reading carries its round's ``version``."""
    out = {}
    if not os.path.isdir(reading_dir):
        return out
    files = [(r, n) for r in sorted(os.listdir(reading_dir)) if os.path.isdir(os.path.join(reading_dir, r))
             for n in sorted(os.listdir(os.path.join(reading_dir, r))) if n.startswith("read_") and n.endswith(".jsonl")]
    for round_name, name in files:
        folder = os.path.join(reading_dir, round_name)
        batch = os.path.join(folder, name.replace("read_", "batch_"))
        path = os.path.join(folder, name)
        if not os.path.exists(batch) or check_reading(path, batch):
            print(f"skipped {round_name}/{name}: run the check on it", file=sys.stderr)
            continue
        version = round_version(folder)
        for line in open(path, encoding="utf-8"):
            if line.strip():
                r = json.loads(line)
                if version >= out.get(r["lead_id"], {}).get("version", 0):
                    out[r["lead_id"]] = {**r, "version": version}
    return out


def priority(r: dict) -> tuple:
    """Call order: payment checks, then ready to close, then the best chance in the next 14 days."""
    return (OUTLOOK_ORDER.get(r.get("outlook"), len(OUTLOOKS)), -(r.get("chance_14d") or 0), r["days_open"])


def _write_csv(path: str, rows: list[dict]) -> None:
    if not rows:
        return
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def _pct_range(r: dict, key: str) -> str:
    return f"{100 * r[key]:.0f}% ({100 * r[key + '_low']:.0f}-{100 * r[key + '_high']:.0f}%)" if r[key] is not None else ""


CALL_LIST = (("Caller", "caller"), ("Team", "team"), ("Kind", "kind"), ("Course", "course"), ("Booked (IST)", "booked_ist"),
             ("Days open", "days_open"), ("Stage", "stage"), ("Days since last real talk", "days_since_spoke"),
             ("Outlook (from the notes)", "outlook"), ("Blocker", "blocker"), ("Situation", "situation"),
             ("Next step", "next_action"), ("What to say", "what_to_say"), ("When", "when"),
             ("Chance in 14 days (from history)", "chance_14"), ("Chance in 45 days", "chance_45"),
             ("Problems seen", "flags"), ("Handling issues", "process_issues"), ("Course fee on lead", "course_fee"),
             ("LeadSquared lead ID", "lead_id"))


def call_list_rows(rows: list[dict]) -> list[dict]:
    out = []
    for r in sorted(rows, key=lambda r: (r["team"], r["caller"], priority(r))):
        out.append({label: ("; ".join(r.get(key) or []) if key in ("flags", "process_issues")
                            else _pct_range(r, "chance_14d") if key == "chance_14"
                            else _pct_range(r, "chance_45d") if key == "chance_45" else r.get(key, ""))
                    for label, key in CALL_LIST})
    return out


def write_workbook(path: str, report: dict, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"Collection audit: open leads booked since {report['since']}. Data as of {report['data_as_of_ist']} IST."])
    ws.append(["Collected = stage reached Collections done or Course Enrolled; rupees wait for payment data. "
               "Chances are the share of past leads in the same position that were collected (90% range), not promises."])
    for title, key in (("By team", "by_kind_team"), ("By caller", "by_kind_caller")):
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        table = report[key]
        if table:
            ws.append(list(table[0]))
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
            for t in table:
                ws.append(list(t.values()))
    for team in sorted({r["team"] for r in rows}):
        sheet = wb.create_sheet(("Calls - " + team)[:31])
        lines = [x for x in call_list_rows(rows) if x["Team"] == team]
        sheet.append([label for label, _ in CALL_LIST])
        for c in sheet[1]:
            c.font = Font(bold=True)
        for x in lines:
            sheet.append([x[label] for label, _ in CALL_LIST])
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for col, (label, _) in zip("ABCDEFGHIJKLMNOPQRST", CALL_LIST):
            sheet.column_dimensions[col].width = 48 if label in ("Situation", "Next step", "What to say", "Problems seen") else 16
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def _counts(rows: list[dict], key: str) -> dict[str, dict]:
    g: dict[str, Counter] = defaultdict(Counter)
    for r in rows:
        if r.get(key):
            g[f'{r["kind"]} / {r["team"]}'][r[key]] += 1
    return {k: dict(c.most_common()) for k, c in sorted(g.items())}


def load_transcripts(path: str, rows: list[dict], pool: list[dict]) -> dict[str, list]:
    """lead_id -> transcript excerpts since its booking, from ``scripts/fetch_collection_transcripts.py``."""
    booked = {v["lead_id"]: v["booked"] for v in pool}
    out = {}
    for line in open(path, encoding="utf-8"):
        d = json.loads(line)
        if d["lead_id"] in booked:
            out[d["lead_id"]] = transcript_excerpts(d["calls"], booked[d["lead_id"]])
    return {r["lead_id"]: out.get(r["lead_id"], []) for r in rows}


def main(data_dir: str, out_dir: str, since: str, transcripts_path: str | None = None) -> None:
    meta, leads, hist = _load(data_dir)
    now = utc(meta["fetched_at_utc"])
    callers = meta.get("team_of_caller") or frozenset(meta.get("team_callers") or ())
    rows, pool, table, views = audit(leads, hist, meta["team_of_owner"], callers, now, since)
    os.makedirs(out_dir, exist_ok=True)
    as_of = now.astimezone(IST).strftime("%Y-%m-%d %H:%M")
    reading_dir = os.path.join(out_dir, "reading")
    readings = load_readings(reading_dir)
    for r in rows:
        r.update({k: v for k, v in readings.get(r["lead_id"], {}).items() if k not in ("lead_id", "version")})
    need = 2 if transcripts_path else 1
    unread = [r for r in rows if readings.get(r["lead_id"], {}).get("version", 0) < need]
    if unread:
        excerpts = load_transcripts(transcripts_path, unread, pool) if transcripts_path else None
        folder = write_reading_round(unread, reading_dir, as_of, excerpts)
        print(f"{len(unread)} open leads still to read: batches and INSTRUCTIONS.md in {folder}", file=sys.stderr)
    _write_csv(os.path.join(out_dir, "open_leads.csv"),
               [{**{k: v for k, v in r.items() if k != "dossier"}, "flags": "; ".join(r["flags"]),
                 "process_issues": "; ".join(r.get("process_issues") or [])} for r in rows])
    report = {
        "data_as_of_ist": as_of, "since": since, "open_leads": len(rows), "read": len(rows) - len(unread),
        "by_kind": pipeline(rows, pool, "kind"), "by_kind_team": pipeline(rows, pool, "kind", "team"),
        "by_kind_caller": pipeline(rows, pool, "kind", "team", "caller"),
        "month_end_ist": month_end(now).strftime("%Y-%m-%d %H:%M"), "new_bookings": new_bookings_outlook(views, now),
        "flags": dict(Counter(f.split(" (")[0].split(" for ")[0] for r in rows for f in r["flags"]).most_common()),
        "outlook": _counts(rows, "outlook"), "blocker": _counts(rows, "blocker"),
        "process_issues": dict(Counter(i for r in rows for i in r.get("process_issues") or []).most_common()),
        "history": [{"key": " / ".join(k), "snapshots": v[0],
                     **{f"collected_{h}d_%": round(100 * x / v[0], 1) for h, x in zip(HORIZONS, v[1:])}}
                    for k, v in sorted(table.items()) if len(k) == 3],
    }
    json.dump(report, open(os.path.join(out_dir, "report.json"), "w"), indent=1, default=str)
    for name in ("by_kind_team", "by_kind_caller"):
        _write_csv(os.path.join(out_dir, f"{name}.csv"), report[name])
    write_workbook(os.path.join(out_dir, "collection_audit.xlsx"), report, rows)
    print(json.dumps({k: report[k] for k in ("data_as_of_ist", "since", "by_kind_team", "flags")}, indent=1))


if __name__ == "__main__":
    if sys.argv[1:2] == ["check"]:
        found = check_reading(sys.argv[2], sys.argv[3])
        print("\n".join(found) or "ok")
        sys.exit(1 if found else 0)
    main(*sys.argv[1:5])
