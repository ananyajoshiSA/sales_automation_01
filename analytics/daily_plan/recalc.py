"""Recalculate a workbook in LibreOffice and count formula errors (the plan must have zero)."""

from __future__ import annotations

import os
import subprocess
import tempfile

from openpyxl import load_workbook


def check(path: str) -> tuple[int, int]:
    """(formulas, errors). An error is a cell LibreOffice left empty or set to '#…' / 'Err:…'."""
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, os.path.abspath(path)],
                       check=True, capture_output=True, timeout=300)
        calc = load_workbook(os.path.join(tmp, os.path.basename(path)), data_only=True)
        src = load_workbook(path)
        n = errors = 0
        for ws in src.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        n += 1
                        v = calc[ws.title][c.coordinate].value
                        if v is None or (isinstance(v, str) and (v.startswith("#") or "Err" in v)):
                            errors += 1
    return n, errors
