"""Revival of older leads: the team's leads from a past period, their whole call history read lead by lead, ranked
by how likely a conversation can be picked up again and turned into an enrollment.

    python -m analytics.daily_plan revival-prepare 2026-10-13 --start 2026-07-01 --end 2026-09-20
    #   ... Claude reads every batch -> data/daily/<date>/revival/reads/out_*.jsonl ...
    python -m analytics.daily_plan revival-build 2026-10-13 [--publish-state]

LeadSquared is only read. Each lead's history comes from one request (all its activities, newest first), so the
period's calls are found without reading every company-wide call since July."""

from __future__ import annotations

import html
import json
import os
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

from analytics.daily_plan.common import CLOSED, REAL_SECS, Snap, ist
from integrations.leadsquared import LeadSquaredClient, parse_activity_note, parse_phone_call

SKIP_STAGES = CLOSED + ("Irrelevant lead", "Invalid")
CALLS, ZIP = (21, 22), 237
RATE = 5  # LeadSquared requests per second (its limit is 30 in 5 seconds)


def pool(snap: Snap, start: str, end: str) -> list[str]:
    """Team-owned leads created in [start, end] (IST days) that are not enrolled, irrelevant or invalid."""
    team = set(snap.users.values())
    out = []
    for lid, l in snap.leads.items():
        t = ist(l.get("CreatedOn"))
        if l.get("OwnerIdName") in team and t and start <= t.strftime("%Y-%m-%d") <= end and l.get("ProspectStage") not in SKIP_STAGES:
            out.append(lid)
    return sorted(out)


def _flat(a: dict) -> dict:
    """A ProspectActivity.svc/Retrieve row in the shape RetrieveByActivityEvent returns (what the parsers expect)."""
    f = a.get("ActivityFields") or {}
    return {**f, "ProspectActivityId": a.get("Id"), "RelatedProspectId": a.get("RelatedProspectId"),
            "ActivityEvent": a.get("EventCode"), "CreatedOn": a.get("CreatedOn"), "ModifiedOn": a.get("ModifiedOn"),
            "ActivityEvent_Note": f.get("ActivityEvent_Note")}


def history(lead_ids: list[str], cache: str, workers: int = 6, client: LeadSquaredClient | None = None) -> dict:
    """{lead_id: [activity rows]} for every lead, one request each; resumable through the JSONL ``cache``."""
    done = {}
    if os.path.exists(cache):
        for line in open(cache):
            try:
                r = json.loads(line)
            except ValueError:  # a line cut short when an earlier run was stopped: read that lead again
                continue
            done[r["lead_id"]] = r["acts"]
    todo = [lid for lid in lead_ids if lid not in done]
    c = client or LeadSquaredClient() if todo else None
    lock, last = threading.Lock(), [0.0]

    def one(lid):
        for attempt in range(5):
            with lock:  # LeadSquared allows 30 calls in 5 seconds: stay under it whatever the thread count
                wait = last[0] + 1 / RATE - time.monotonic()
                if wait > 0:
                    time.sleep(wait)
                last[0] = time.monotonic()
            try:
                r = c.get_lead_activities(lid, row_count=200) or {}
                return lid, [_flat(a) for a in r.get("ProspectActivities") or [] if a.get("EventCode") in CALLS + (ZIP,)]
            except Exception as e:  # noqa: BLE001 - a 429 or a dropped connection: back off and try again
                if attempt == 4:
                    raise
                print(f"history: {lid[:8]} retry after {type(e).__name__}", file=sys.stderr, flush=True)
                time.sleep(5 * (attempt + 1))

    if os.path.exists(cache) and os.path.getsize(cache) and open(cache, "rb").read()[-1:] != b"\n":
        open(cache, "a").write("\n")  # never glue a new line onto one cut short
    with open(cache, "a") as out, ThreadPoolExecutor(workers) as ex:
        for k, (lid, acts) in enumerate(ex.map(one, todo), 1):
            done[lid] = acts
            out.write(json.dumps({"lead_id": lid, "acts": acts}) + "\n")
            out.flush()
            if k % 200 == 0:
                print(f"history: {k} of {len(todo)} leads read", file=sys.stderr, flush=True)
    return done


def snapshot(base: dict, hist: dict) -> dict:
    """A team snapshot whose calls and Zip notes are the leads' full histories (for Snap and the dossier builder)."""
    calls, zips = [], []
    for lid, acts in hist.items():
        for a in acts:
            if a["ActivityEvent"] in CALLS:
                calls.append(parse_phone_call(a))
            else:
                zips.append({**a, "note": parse_activity_note(a.get("ActivityEvent_Note"))})
    return {**base, "calls": calls, "zip_activities": zips, "payments": []}


def shortlist(snap: Snap, lead_ids: list[str]) -> list[str]:
    """Leads that had at least one real conversation (answered, 2 min+): counselling happened."""
    out = []
    for lid in lead_ids:
        cs = snap.by_lead.get(lid, [])
        if any(c["status"] == "Answered" and (c.get("duration") or 0) >= REAL_SECS for c in cs):
            out.append(lid)
    return out


def focus(snap: Snap, lead_ids: list[str], since: str, ni_min_secs: int = 600) -> tuple[list[str], dict]:
    """Leads to read: counselled, no real conversation since ``since`` (those are in the daily reports already), and,
    for leads marked Not Interested, a call of at least ``ni_min_secs`` (long counselling is where wrong tags hide)."""
    picked, skipped = [], {"no real conversation": 0, "real conversation since " + since + " (in the daily reports)": 0,
                           "Not Interested after short calls only": 0}
    for lid in lead_ids:
        real = [c for c in snap.by_lead.get(lid, []) if c["status"] == "Answered" and (c.get("duration") or 0) >= REAL_SECS]
        if not real:
            skipped["no real conversation"] += 1
        elif any(c["t"].strftime("%Y-%m-%d") >= since for c in real):
            skipped["real conversation since " + since + " (in the daily reports)"] += 1
        elif (snap.leads.get(lid) or {}).get("ProspectStage") == "Not Interested" and max(c["duration"] for c in real) < ni_min_secs:
            skipped["Not Interested after short calls only"] += 1
        else:
            picked.append(lid)
    return picked, skipped


def numbers(snap: Snap, lead_ids: list[str]) -> list[str]:
    from integrations.transcripts import normalize_phone

    seen, out = set(), []
    for lid in lead_ids:
        n = normalize_phone((snap.leads.get(lid) or {}).get("Phone") or next((snap.phone(c) for c in snap.by_lead.get(lid, []) if snap.phone(c)), ""))
        if n and n not in seen:
            seen.add(n)
            out.append(n)
    return out


def load(reads_dir: str) -> list[dict]:
    rows = []
    for f in sorted(os.listdir(reads_dir)):
        if f.startswith("out_") and f.endswith(".jsonl"):
            rows += [json.loads(l) for l in open(os.path.join(reads_dir, f)) if l.strip()]
    return rows


BANDS = (("High", 20), ("Medium", 10), ("Low", 5))


def band(chance: int) -> str:
    return next((name for name, lo in BANDS if chance >= lo), "Reject")


def ranked(rows: list[dict]) -> list[dict]:
    """Leads worth a call, highest chance first; ties go to the most recent real conversation."""
    keep = [r for r in rows if r.get("revive") and int(r.get("revival_chance") or 0) >= BANDS[-1][1]]
    keep = sorted(keep, key=lambda r: r.get("last_real_talk_date") or "", reverse=True)
    return sorted(keep, key=lambda r: -int(r.get("revival_chance") or 0))


def batches(index: list[dict], max_leads: int = 30, max_bytes: int = 220_000) -> list[list[dict]]:
    """Reading batches of at most ``max_leads`` dossiers and about ``max_bytes`` of text, one owner's leads together."""
    out, size = [[]], 0
    for i in sorted(index, key=lambda i: (i["owner"] or "", i["lead_id"])):
        if out[-1] and (len(out[-1]) >= max_leads or size + i["size"] > max_bytes):
            out.append([])
            size = 0
        out[-1].append(i)
        size += i["size"]
    return [b for b in out if b]


def to_state(date: str, rows: list[dict]) -> dict:
    from analytics.daily_plan.state import h

    return {"date": date, "kind": "revival",
            "leads": {h(r["lead_id"]): {"owner": r.get("owner"), "chance": int(r.get("revival_chance") or 0),
                                        "band": band(int(r.get("revival_chance") or 0))} for r in ranked(rows)}}


CSS = """@page { size: A4 landscape; margin: 10mm; } body { font-family: Arial, sans-serif; font-size: 9px; color: #1d2b28; }
h1 { color: #1F4E46; font-size: 20px; margin: 0 0 4px; } h2 { color: #1F4E46; font-size: 14px; margin: 14px 0 6px; page-break-after: avoid; }
table { border-collapse: collapse; width: 100%; } th { background: #E8EFEC; text-align: left; } th, td { border: 1px solid #d6dcda; padding: 3px 4px; vertical-align: top; }
tr { page-break-inside: avoid; } .box { border-left: 4px solid #1F4E46; background: #F4F7F6; padding: 8px 10px; margin: 6px 0; font-size: 10px; }
.small { color: #555; font-size: 8.5px; } .High { background: #DCEFE6; } .Medium { background: #EEF6F1; }"""


def html_doc(rows: list[dict], meta: dict, team: str) -> str:
    e = html.escape
    keep = ranked(rows)
    rej = [r for r in rows if r not in keep]
    by_band = {b: [r for r in keep if band(int(r.get("revival_chance") or 0)) == b] for b, _ in BANDS}
    owners = sorted({r.get("owner") or "?" for r in keep})
    H = [f"<!doctype html><html><head><meta charset='utf-8'><title>Revival list</title><style>{CSS}</style></head><body>",
         f"<h1>{e(team)}: leads from {e(meta['period'])} that can be revived</h1>",
         f"<p class='small'>Built {e(meta['built'])} IST from each lead's full LeadSquared call history, Zipteams notes and call transcripts, "
         "every lead read by Claude. Chances are estimates of an enrollment within 30 days of a fresh call; they are not targets.</p>",
         f"<div class='box'><b>{meta['pool']:,}</b> leads were created in the period and are not enrolled, irrelevant or invalid. "
         f"<b>{meta['talked']:,}</b> were counselled and left untouched since 21 Sep, and were read in full. "
         f"<b>{len(keep)}</b> can be revived: {len(by_band['High'])} high (20%+), {len(by_band['Medium'])} medium (10–19%), "
         f"{len(by_band['Low'])} low (5–9%). {len(rej)} were rejected. Expected enrollments if every one is called "
         f"(estimate): about {sum(int(r.get('revival_chance') or 0) for r in keep) / 100:.0f}.</div>",
         ("<p class='small'>Not read: " + "; ".join(f"{v:,} {k}" for k, v in meta.get("skipped", {}).items()) + ".</p>") if meta.get("skipped") else "",
         "<h2>By caller</h2><table><tr><th>Caller</th><th>High</th><th>Medium</th><th>Low</th><th>Total</th></tr>"]
    for o in owners:
        n = [sum(1 for r in by_band[b] if (r.get("owner") or "?") == o) for b, _ in BANDS]
        H.append(f"<tr><td>{e(o)}</td>" + "".join(f"<td>{x}</td>" for x in n) + f"<td>{sum(n)}</td></tr>")
    H.append("</table>")
    reasons = {}
    for r in rej:
        reasons[r.get("reject_reason") or "low chance"] = reasons.get(r.get("reject_reason") or "low chance", 0) + 1
    H.append("<h2>Why leads were rejected</h2><table><tr><th>Reason</th><th>Leads</th></tr>"
             + "".join(f"<tr><td>{e(k)}</td><td>{v}</td></tr>" for k, v in sorted(reasons.items(), key=lambda kv: -kv[1])[:15]) + "</table>")
    fix = [(k, [r for r in rej if r.get("reject_reason") == k]) for k in ("already enrolled or paid", "existing student", "do not call")]
    if any(rs for _, rs in fix):
        H.append("<h2>Records to fix in LeadSquared (not sales calls)</h2><table><tr><th>What</th><th>Leads</th></tr>")
        for k, rs in fix:
            if rs:
                H.append(f"<tr><td>{e(k)} ({len(rs)})</td><td>" + "; ".join(
                    f"{e(r.get('name') or '?')} ({e((r.get('owner') or '?').split()[0])})" for r in rs[:60]) + "</td></tr>")
        H.append("</table><p class='small'>Paid or existing students: move to Course Enrolled or send to support. "
                 "Do not call: mark the lead so neither callers nor the AI campaigns dial it.</p>")
    H.append("<h2 style='page-break-before:always'>The revival list, most likely first</h2>"
             "<table><tr><th style='width:2%'>#</th><th style='width:11%'>Lead · caller</th><th style='width:9%'>Course · fee</th>"
             "<th style='width:7%'>Chance · stage reached</th><th style='width:22%'>What happened</th><th style='width:13%'>Why it stalled</th>"
             "<th style='width:14%'>Hook to reopen</th><th style='width:14%'>Opening line</th><th style='width:8%'>Next step</th></tr>")
    for i, r in enumerate(keep, 1):
        ch = int(r.get("revival_chance") or 0)
        H.append(f"<tr class='{band(ch)}'><td>{i}</td><td><b>{e(r.get('name') or '')}</b><br>{e(r.get('owner') or '')}<br>"
                 f"<span class='small'>{e(r.get('phone') or '')}</span></td>"
                 f"<td>{e(r.get('course') or '')}<br><span class='small'>{e(r.get('fee_discussed') or '')}</span></td>"
                 f"<td><b>{ch}%</b><br>{e(r.get('stage_reached') or '')}<br><span class='small'>last real talk {e(r.get('last_real_talk') or '')}</span></td>"
                 f"<td>{e(r.get('summary') or '')}</td><td>{e(r.get('why_stalled') or '')}</td><td>{e(r.get('revival_hook') or '')}</td>"
                 f"<td>{e(r.get('opening_line') or '')}</td><td>{e(r.get('next_step') or '')}</td></tr>")
    H.append("</table></body></html>")
    return "\n".join(H)


def workbook(rows: list[dict], out: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    keep = ranked(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Revival list"
    head = ["#", "Band", "Chance %", "Lead", "Phone", "Caller", "Course", "Fee discussed", "Stage reached", "Last real talk",
            "What happened", "Why it stalled", "Hook to reopen", "Opening line", "Next step", "Outcome", "Note"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E46")
    fills = {"High": "DCEFE6", "Medium": "EEF6F1"}
    for i, r in enumerate(keep, 1):
        ch = int(r.get("revival_chance") or 0)
        ws.append([i, band(ch), ch, r.get("name"), r.get("phone"), r.get("owner"), r.get("course"), r.get("fee_discussed"),
                   r.get("stage_reached"), r.get("last_real_talk"), r.get("summary"), r.get("why_stalled"), r.get("revival_hook"),
                   r.get("opening_line"), r.get("next_step"), "", ""])
        if band(ch) in fills:
            for c in ws[ws.max_row][:3]:
                c.fill = PatternFill("solid", fgColor=fills[band(ch)])
    for col, w in zip("ABCDEFGHIJKLMNOPQ", (4, 8, 8, 22, 15, 16, 22, 16, 16, 12, 50, 30, 30, 34, 26, 16, 20)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    rj = wb.create_sheet("Rejected")
    rj.append(["Lead", "Caller", "Reason", "What happened"])
    for r in rows:
        if r not in keep:
            rj.append([r.get("name"), r.get("owner"), r.get("reject_reason") or "low chance", r.get("summary")])
    wb.save(out)
