"""Workbook: team leader sheet, Team summary, How to use, one sheet per caller. Formulas count an enrollment only with a UTR."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
F = "Arial"; HDR = PatternFill("solid", fgColor="1F4E46"); INP = PatternFill("solid", fgColor="FFF4CC")
TF = {"M": "FBE0DE", "P": "FFF0C2", "A": "DCEFE6", "B": "E8F0FA", "F": "F3EAF7", "R": "FBEFE3", "C": "F2F2F2"}
TL = {"P": "P · Paid booking – collect balance", "M": "M · Return the lead's call", "A": "A · Close today", "B": "B · Hot follow-up", "F": "F · New lead, first contact",
      "R": "R · Revive closure-stage", "C": "C · Nurture (WhatsApp first)"}
ATT = {"P": 3, "M": 3, "A": 4, "B": 3, "F": 3, "R": 2, "C": 1}
S = Side(style="thin", color="D0D5D2"); BD = Border(left=S, right=S, top=S, bottom=S)
OUTCOMES = ["Paid – new enrollment", "Already a student", "Payment link sent", "Seat blocked / token paid", "EMI docs pending",
            "Callback fixed", "Spoke – thinking", "Not answered", "Switched off / failed", "Busy – call later", "Not interested", "Wrong number"]
PAY = ["Paid", "Token paid", "Link sent", "EMI docs pending", "Not yet"]
def hdr(ws, row, vals, widths=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v); c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BD
        if widths: ws.column_dimensions[L(i)].width = widths[i - 1]
def put(ws, r, c, v, bold=False, wrap=True, fill=None, size=9, italic=False, color=None):
    x = ws.cell(row=r, column=c, value=v); x.font = Font(name=F, size=size, bold=bold, italic=italic, color=color)
    x.alignment = Alignment(wrap_text=wrap, vertical="top"); x.border = BD
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    return x
def title(ws, r, text, size=12):
    x = ws.cell(row=r, column=1, value=text); x.font = Font(name=F, bold=True, size=size, color="1F4E46")



def write(R: dict, C: dict, out: str) -> str:
    LN = C.get("leader", "Team leader").split()[0]
    PL = C["review"]["day_label"]
    SLOT = C["slots"]
    wb = Workbook(); ws = wb.active; ws.title = C.get("leader", "Team leader").split()[0]
    ws["A1"] = C["shivangi_title"]; ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = C["shivangi_intro"]; ws["A2"].font = Font(name=F, size=10); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:P2"); ws.row_dimensions[2].height = 75
    r = 4; title(ws, r, "1. Your priority leads — where today's enrollments come from. The caller calls; you own the outcome."); r += 1
    PH = ["#", "Caller", "Lead", "Phone", "Course", "Est. chance (3 days)", "Why this lead", "What the caller must do", f"{LN}'s role",
          "Check by", "Who decides", "How they can pay", "Called by check time? (Y/N)", "Outcome", "Payment proof (UTR)", f"{LN}'s notes"]
    hdr(ws, r, PH, [4, 15, 22, 15, 26, 10, 50, 46, 44, 9, 16, 22, 11, 20, 16, 30]); r += 1
    dvo = DataValidation(type="list", formula1='"' + ",".join(OUTCOMES) + '"', allow_blank=True); ws.add_data_validation(dvo)
    dvy = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); ws.add_data_validation(dvy)
    cur = None
    for i, p in enumerate(C["priority"], 1):
        if p.get("group") != cur:
            cur = p.get("group"); n = sum(1 for q in C["priority"] if q.get("group") == cur)
            x = ws.cell(row=r, column=1, value=f"{C['priority_groups'][cur]} ({n})"); x.font = Font(name=F, bold=True, size=10, color="FFFFFF"); x.fill = HDR
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=16); r += 1
        vals = [i, p["owner"], p["name"], "+91-" + str(p["phone"]) if p.get("phone") else "", p.get("course", ""), p["chance"] / 100,
                p["why"], p["ask"], p["role"], p["check_by"], p.get("who_decides", ""), p.get("how_pay", "")]
        for j, v in enumerate(vals, 1): put(ws, r, j, v, bold=j == 3)
        ws.cell(row=r, column=6).number_format = "0%"
        for j in (13, 14, 15, 16): put(ws, r, j, None, fill="FFF4CC")
        dvy.add(f"M{r}"); dvo.add(f"N{r}"); ws.row_dimensions[r].height = 120; r += 1
    r += 1; title(ws, r, "2. Your checks today — what to look at, and what to do if it fails"); r += 1
    hdr(ws, r, ["Time", "Check", "", "", "", "", "If it fails, do this"]); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6); r += 1
    for ck in C["checks"]:
        put(ws, r, 1, ck[0], bold=True); put(ws, r, 2, ck[1]); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        put(ws, r, 7, ck[2]); ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9); ws.row_dimensions[r].height = 48; r += 1
    r += 1; title(ws, r, f"3. Feedback to give each caller (from {PL}'s calls and transcripts)"); r += 1
    hdr(ws, r, ["", "Caller", f"{PL}'s numbers", "", "", "", "What went well", "One thing to fix", "Their 3 rules today"]); r += 1
    for fb in C["feedback"]:
        put(ws, r, 2, fb["caller"], bold=True); put(ws, r, 3, fb["numbers"]); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        put(ws, r, 7, fb["good"]); put(ws, r, 8, fb["fix"]); put(ws, r, 9, "\n".join(f"{k}. {x}" for k, x in enumerate(fb["rules"], 1)))
        ws.row_dimensions[r].height = 130; r += 1
    r += 1; title(ws, r, "4. Team-wide changes from today (dialling pattern)"); r += 1
    for ch in C["changes"]:
        x = ws.cell(row=r, column=1, value=ch); x.font = Font(name=F, size=10); x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12); ws.row_dimensions[r].height = 42; r += 1
    ws.freeze_panes = "A4"

    COLS = [("#", 5), ("Tier", 20), ("When to call", 22), ("Lead", 22), ("Phone", 15), ("Course (confirmed)", 26), ("Est. chance (3 days)", 10),
            ("Missed call not returned", 10), ("Why this lead", 50), ("Opening line", 44), ("Exact ask", 46), ("Prepare for (objection)", 34),
            ("Who decides", 16), ("How they can pay", 20), ("Real blocker", 24), ("Best time (past pick-ups)", 16), ("Call only on / WhatsApp only", 18),
            ("Today so far", 22), ("Last real conversation", 13), ("Stage", 16), ("Attempts required", 9),
            ("Attempt 1 time", 9), ("Attempt 1 outcome", 18), ("Attempt 2 time", 9), ("Attempt 2 outcome", 18), ("Attempt 3 time", 9),
            ("Attempt 3 outcome", 18), ("Attempt 4 time", 9), ("Attempt 4 outcome", 18), ("WhatsApp sent (Y/N)", 9), ("Next step date/time", 15),
            ("Outcome", 20), ("Payment status", 14), ("Payment proof (UTR / txn ID)", 18), ("Attempts done", 9), ("Status", 15), ("Notes", 30), ("Lead ID", 36)]
    COL = {n: i + 1 for i, (n, _) in enumerate(COLS)}
    INPUTS = [n for n, _ in COLS if n.startswith("Attempt ") and "required" not in n] + ["WhatsApp sent (Y/N)", "Next step date/time", "Outcome", "Payment status", "Payment proof (UTR / txn ID)", "Notes"]
    order = C["owner_order"]; ranges = {}
    def caller_sheet(owner, rows):
        ws = wb.create_sheet(owner[:31]); fb = next((f for f in C["feedback"] if f["caller"] == owner), None)
        ws["A1"] = f"{owner} — call plan for {C['day_label']}"; ws["A1"].font = Font(name=F, bold=True, size=14)
        lines = []
        if fb:
            lines = [fb["numbers"], f"Went well: {fb['good']}", f"Fix: {fb['fix']}", "Your 3 rules today: " + "  ".join(f"{k}. {x}" for k, x in enumerate(fb["rules"], 1))]
        lines.append(C["caller_note"])
        for k, ln in enumerate(lines, 2):
            ws.cell(row=k, column=1, value=ln).font = Font(name=F, size=10, bold=ln.startswith("Your 3"))
            ws.cell(row=k, column=1).alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells(start_row=k, start_column=1, end_row=k, end_column=14)
            ws.row_dimensions[k].height = 30 if len(ln) < 160 else 44
        top = len(lines) + 2
        first, last = top + 2, top + 1 + max(len(rows), 1)
        sc = L(COL["Status"]); ws.cell(row=top, column=1, value="Attempts logged").font = Font(name=F, bold=True, size=10)
        ws.cell(row=top, column=3, value=f"=SUM({L(COL['Attempts done'])}{first}:{L(COL['Attempts done'])}{last})")
        ws.cell(row=top, column=4, value="Enrolled (with proof)").font = Font(name=F, bold=True, size=10)
        ws.cell(row=top, column=5, value=f'=COUNTIF({sc}{first}:{sc}{last},"Enrolled")')
        ws.cell(row=top, column=6, value="Paid? add proof").font = Font(name=F, bold=True, size=10)
        ws.cell(row=top, column=7, value=f'=COUNTIF({sc}{first}:{sc}{last},"Paid? add proof")')
        ws.cell(row=top, column=8, value="Expected enrollments (sum of chances, estimate)").font = Font(name=F, bold=True, size=10)
        ws.cell(row=top, column=9, value=f"=SUM({L(COL['Est. chance (3 days)'])}{first}:{L(COL['Est. chance (3 days)'])}{last})").number_format = "0.0"
        hdr(ws, top + 1, [n for n, _ in COLS], [w for _, w in COLS]); ws.row_dimensions[top + 1].height = 42
        dvo = DataValidation(type="list", formula1='"' + ",".join(OUTCOMES) + '"', allow_blank=True)
        dvp = DataValidation(type="list", formula1='"' + ",".join(PAY) + '"', allow_blank=True)
        dvy = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
        for dv in (dvo, dvp, dvy): ws.add_data_validation(dv)
        g = lambda n, rr: f"{L(COL[n])}{rr}"
        for i, r in enumerate(rows):
            rr = first + i; tier = r["tier"]
            vals = {"#": i + 1, "Tier": ("A · VERIFY payment (marked enrolled)" if r.get("verify") else TL[tier]), "When to call": SLOT[tier],
                    "Lead": r.get("name") or "Unnamed lead", "Phone": ("+91-" + r["phone"]) if r.get("phone") else "", "Course (confirmed)": r.get("course", ""),
                    "Est. chance (3 days)": r["chance"] / 100, "Missed call not returned": (f"Yes ({r['missed_last']})" if r.get("missed") else ""),
                    "Why this lead": r.get("why", ""), "Opening line": r.get("opening_line", ""), "Exact ask": r.get("exact_ask", ""),
                    "Prepare for (objection)": r.get("objection", ""), "Who decides": r.get("who_decides", ""), "How they can pay": r.get("how_pay", ""),
                    "Real blocker": r.get("real_blocker", ""), "Best time (past pick-ups)": r.get("best_time", ""),
                    "Call only on / WhatsApp only": " ".join(x for x in [r.get("callback_requested") or "", "WhatsApp only" if r.get("whatsapp_only") else ""] if x),
                    "Today so far": r.get("today", ""), "Last real conversation": r.get("last_conv", ""), "Stage": r.get("stage", ""),
                    "Attempts required": ATT[tier], "Lead ID": r["lead_id"]}
            for n, v in vals.items():
                x = put(ws, rr, COL[n], v, bold=n in ("Lead", "Tier"), fill=TF[tier])
            ws.cell(row=rr, column=COL["Est. chance (3 days)"]).number_format = "0%"
            ws.cell(row=rr, column=COL["Attempts done"], value="=COUNTA(" + ",".join(g(f"Attempt {k} time", rr) for k in range(1, 5)) + ")")
            enr = f'OR({g("Outcome", rr)}="Paid – new enrollment",' + ",".join(f'{g(f"Attempt {k} outcome", rr)}="Paid – new enrollment"' for k in range(1, 5)) + f',{g("Payment status", rr)}="Paid")'
            stu = f'OR({g("Outcome", rr)}="Already a student",' + ",".join(f'{g(f"Attempt {k} outcome", rr)}="Already a student"' for k in range(1, 5)) + ")"
            ws.cell(row=rr, column=COL["Status"], value=(
                f'=IF({stu},"Already a student",IF({enr},IF(LEN(TRIM({g("Payment proof (UTR / txn ID)", rr)}))>0,"Enrolled","Paid? add proof"),'
                f'IF({g("Attempts done", rr)}>={g("Attempts required", rr)},"Done",IF({g("Attempts done", rr)}=0,"Not started","In progress"))))'))
            for n in ("Attempts done", "Status"): put(ws, rr, COL[n], ws.cell(row=rr, column=COL[n]).value, bold=True)
            for n in INPUTS: x = ws.cell(row=rr, column=COL[n]); x.fill = INP; x.border = BD; x.font = Font(name=F, size=9)
            for k in range(1, 5): dvo.add(g(f"Attempt {k} outcome", rr))
            dvo.add(g("Outcome", rr)); dvp.add(g("Payment status", rr)); dvy.add(g("WhatsApp sent (Y/N)", rr))
            ws.row_dimensions[rr].height = 96 if tier in "MAB" else 48
        rng = f"{sc}{first}:{sc}{last}"
        for text, col in (("Enrolled", "B7E1C1"), ("Paid? add proof", "FFD966"), ("Already a student", "D9D2E9"), ("Not started", "F8D7D3"), ("Done", "E2E2E2")):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{sc}{first}="{text}"'], fill=PatternFill("solid", fgColor=col)))
        ws.freeze_panes = ws.cell(row=first, column=COL["Phone"]); ws.auto_filter.ref = f"A{top+1}:{L(len(COLS))}{last}"
        ranges[owner] = (ws.title, first, last)

    summ = wb.create_sheet("Team summary", 1); guide = wb.create_sheet("How to use", 2)
    for o in order: caller_sheet(o, R["by_owner"].get(o, []))
    summ["A1"] = f"Team summary — {C['day_label']} (updates as callers fill their sheets)"; summ["A1"].font = Font(name=F, bold=True, size=14)
    H = ["Caller", "Leads", "M", "P", "A", "B", "F", "R", "C", "Attempts logged", "Not started (3+ attempts due)", "Enrolled (with proof)", "Paid? add proof", "Already a student", "Expected enrollments (estimate)", "Target"]
    hdr(summ, 3, H, [20, 8, 6, 6, 6, 6, 6, 6, 6, 11, 12, 12, 12, 12, 14, 8])
    for i, o in enumerate(order):
        rr = 4 + i; sh, f, l = ranges[o]; q = f"'{sh}'!"
        T = lambda col: f"{q}${L(COL[col])}${f}:${L(COL[col])}${l}"
        rows = R["by_owner"].get(o, [])
        vals = [o, len(rows)] + [sum(r["tier"] == t for r in rows) for t in "MPABFRC"]
        for j, v in enumerate(vals, 1): put(summ, rr, j, v, bold=j == 1, size=10)
        put(summ, rr, 10, f"=SUM({T('Attempts done')})", size=10)
        put(summ, rr, 11, f'=COUNTIFS({T("Status")},"Not started",{T("Attempts required")},">=3")', size=10)
        put(summ, rr, 12, f'=COUNTIF({T("Status")},"Enrolled")', size=10, bold=True)
        put(summ, rr, 13, f'=COUNTIF({T("Status")},"Paid? add proof")', size=10)
        put(summ, rr, 14, f'=COUNTIF({T("Status")},"Already a student")', size=10)
        put(summ, rr, 15, f"=SUM({T('Est. chance (3 days)')})", size=10).number_format = "0.0"
        put(summ, rr, 16, C["targets"].get(o, 4), size=10)
    tr = 4 + len(order)
    put(summ, tr, 1, "Team", bold=True, size=10)
    for j in range(2, 17):
        if j == 16: continue
        c = put(summ, tr, j, f"=SUM({L(j)}4:{L(j)}{tr-1})", bold=True, size=10)
        if j == 15: c.number_format = "0.0"
    summ.cell(row=tr + 2, column=1, value="Enrolled counts only rows with Outcome 'Paid – new enrollment' (or Payment status 'Paid') AND a UTR/transaction ID in Payment proof. 'Already a student' is never counted. Expected enrollments are judgement-based estimates, not targets.").font = Font(name=F, italic=True, size=9)
    guide["A1"] = "How to use this workbook"; guide["A1"].font = Font(name=F, bold=True, size=14); guide.column_dimensions["A"].width = 130
    for k, ln in enumerate(C["how_to"], 3):
        x = guide.cell(row=k, column=1, value=ln); x.font = Font(name=F, size=10, bold=ln.endswith(":")); x.alignment = Alignment(wrap_text=True)
    wb.save(out)
    return out
