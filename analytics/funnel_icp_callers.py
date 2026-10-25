"""Stage 2 of the funnel ICP report: one row per lead for human sales callers.

Each lead gets the ICP to use (what they told the AI agent, when known, else the ad they came from), what
they said on the AI call, a priority, and the opening line, story and bootcamp day for their ICP from the
map's ``playbook``. Leads who already bought go to their own tab; leads who refused, gave a wrong number
or were blacklisted go to "Do not call".

Priority (no invented probabilities; the ICP rate is last measured lead → sale for that ICP):
    1  Positive on the AI call, and their ICP converts at or above the campaign average
    2  Positive on the AI call (other ICPs), or asked for a callback
    3  Spoke to the AI agent, nothing positive yet
    4  Never connected

Written by ``python -m analytics.funnel_icp report ...`` next to the stage 1 file; ``--stated-icp`` adds
the ICP read from each lead's call summary (a TSV of phone, code).
"""

from __future__ import annotations

from typing import Any

from analytics.funnel_icp import POSITIVE, first_touch, page_slug, person_key, phone10, email_key

DO_NOT_CALL = ("Not Interested", "Wrong Number", "Blacklisted", "Irrelevant", "Busy – No Callback")
CALLBACK = ("Call Back", "Third Party Callback")
PRIORITY_TEXT = {1: "1 · Call first", 2: "2 · Call today", 3: "3 · Call if time", 4: "4 · Not reached yet"}


def read_stated(path: str | None, codes: dict[str, str]) -> dict[str, str]:
    """phone -> ICP from a TSV of ``phone<TAB>code`` (codes as in the map's ``stated_icp_codes``)."""
    out: dict[str, str] = {}
    if not path:
        return out
    with open(path, encoding="utf-8") as f:
        for line in f:
            parts = line.rstrip("\n").split("\t")
            if len(parts) >= 2 and parts[1].strip() in codes:
                out[phone10(parts[0])] = codes[parts[1].strip()]
    return out


def caller_rows(leads: list[dict], result: dict, icp_map: dict, tp: dict[str, dict] | None,
                stated: dict[str, str]) -> list[dict]:
    pages, segs, book = icp_map["pages"], icp_map.get("general_segments", {}), icp_map.get("playbook", {})
    people = first_touch(leads)
    sold = {s["key"]: s for s in result["sales"] if s["key"]}
    total_leads = sum(v["leads"] for v in result["icp"].values()) or 1
    avg = sum(v["sales"] for v in result["icp"].values()) / total_leads
    rate = {k: (v["sales"] / v["leads"] if v["leads"] else 0.0) for k, v in result["icp"].items()}
    rows = []
    for key, row in people.items():
        slug = page_slug(row.get("pageUrl"))
        ad_icp = pages.get(slug, "UNMAPPED")
        said = stated.get(key, "")
        use = said or ad_icp
        # "General" from the call only means none of the call ICPs fit, so it never replaces a specific ad
        # ICP; teacher ads keep Teachers when the reading had no teacher option (Retired / General / Senior women).
        if said and ad_icp not in ("General", "UNMAPPED") and (said == "General" or (
                ad_icp == "Teachers / educators" and said in ("Retired / about to retire", "Senior women professionals"))):
            use = ad_icp
        t = (tp or {}).get(key) or {}
        dispos = t.get("dispositions", [])
        positive = t.get("best", 99) < len(POSITIVE)
        if key in sold:
            tab = "Already enrolled"
        elif any(d.startswith(DO_NOT_CALL) for d in dispos) and not positive:
            tab = "Do not call"
        else:
            tab = "Call list"
        if positive and rate.get(use, 0) >= avg:
            prio, why = 1, f"Said yes to the AI agent; {use} converts above average"
        elif positive:
            prio, why = 2, "Said yes to the AI agent"
        elif any(d.startswith(CALLBACK) for d in dispos):
            prio, why = 2, "Asked for a callback"
        elif t.get("connected"):
            prio, why = 3, "Spoke to the AI agent, no yes yet"
        else:
            prio, why = 4, "Never connected" if t else "Not dialled by the AI agent"
        pb = book.get(use, {})
        rows.append({
            "tab": tab, "prio": prio, "Priority": PRIORITY_TEXT[prio], "Why": why,
            "Name": row.get("name") or row.get("fullName") or "", "Phone": key if key.isdigit() else "",
            "Email": email_key(row.get("email")), "ICP to use": use, "ICP from ad": ad_icp,
            "ICP from what they told the AI": said or "(not read yet)",
            "Check": "Ad and call disagree: using what they told the AI" if said and use != ad_icp else "",
            "Segment": segs.get(slug, "") if ad_icp == "General" else "",
            "Ad they responded to": row.get("funnel") or "", "Signed up (IST)": row.get("captured_ist") or "",
            "AI call result": t.get("disposition") or ("No answer" if t else "Not dialled"),
            "AI attempts": t.get("attempts", 0), "Last AI call (IST)": t.get("last_call", ""),
            "What they told the AI agent": t.get("summary", ""),
            "Pain to touch": pb.get("pain", ""), "Open with": pb.get("open_with", ""),
            "Story to tell": pb.get("story", ""), "Day to point to": pb.get("day", ""),
            "ICP lead → sale (last measured)": round(rate.get(use, 0), 4),
            "Enrolled status": sold[key]["status"] if key in sold else "",
            "Call outcome (fill in)": "", "Next step (fill in)": "",
        })
    for r in rows:
        if r["tab"] == "Already enrolled":
            r["Priority"], r["Why"] = "", "Already enrolled: hand over to onboarding"
        elif r["tab"] == "Do not call":
            r["Priority"], r["Why"] = "", "Refused, wrong number or blacklisted on the AI call"
    rows.sort(key=lambda r: (r["prio"], -r["ICP lead → sale (last measured)"], r["Last AI call (IST)"] or "~"))
    return rows


COLS = ["Priority", "Why", "Name", "Phone", "Email", "ICP to use", "Check", "Open with", "Story to tell",
        "Day to point to", "Pain to touch", "What they told the AI agent", "AI call result", "AI attempts",
        "Last AI call (IST)", "ICP from ad", "ICP from what they told the AI", "Segment", "Ad they responded to",
        "Signed up (IST)", "ICP lead → sale (last measured)", "Call outcome (fill in)", "Next step (fill in)"]


def caller_checks(rows: list[dict], result: dict) -> list[tuple[str, str, str]]:
    out = []
    n = len(rows)
    out.append(("PASS" if n == result["people"] else "FAIL", "Caller list: one row per lead",
                f"{n:,} rows vs {result['people']:,} leads"))
    tabs = {r["tab"] for r in rows}
    out.append(("PASS" if tabs <= {"Call list", "Already enrolled", "Do not call"} else "FAIL",
                "Caller list: every lead in exactly one tab", ", ".join(sorted(tabs))))
    enrolled = sum(1 for r in rows if r["tab"] == "Already enrolled")
    matched = sum(1 for s in result["sales"] if s["key"])
    keys = len({s["key"] for s in result["sales"] if s["key"]})
    out.append(("PASS" if enrolled == keys else "FAIL", "Caller list: enrolled tab = matched sales (people)",
                f"{enrolled} vs {keys} people ({matched} sale rows)"))
    by_ad: dict[str, int] = {}
    for r in rows:
        by_ad[r["ICP from ad"]] = by_ad.get(r["ICP from ad"], 0) + 1
    bad = {k: (v, result["icp"].get(k, {}).get("leads")) for k, v in by_ad.items()
           if v != result["icp"].get(k, {}).get("leads")}
    out.append(("PASS" if not bad else "FAIL", "Caller list: leads per ad ICP = stage 1", str(bad) if bad else "ok"))
    return out


def write_caller_xlsx(rows: list[dict], path: str, title: str, fresh: str, icp_map: dict, result: dict) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "How to use"
    for line in [
        [title], [fresh], [],
        ["Who this is for: sales callers. One row per lead, best leads first."],
        ["1. Work the Call list tab top to bottom. Priority 1 first, then 2, 3, 4."],
        ["2. Before dialling, read 'What they told the AI agent'. Don't ask again what they already told us."],
        ["3. Use 'Open with', 'Story to tell' and 'Day to point to' for their ICP. Never say the word ICP."],
        ["4. 'ICP to use' is what they told our AI agent when we know it; otherwise the ad they came from."],
        ["   If the Check column says the ad and call disagree, trust the call."],
        ["5. Fill in 'Call outcome' and 'Next step' after every call."],
        ["6. Do not call anyone on the 'Do not call' tab. 'Already enrolled' is for onboarding, not sales."],
        [],
        ["Priority", "Meaning"],
        *[[PRIORITY_TEXT[k], v] for k, v in {
            1: "Said yes to the AI agent, and their ICP buys at or above the campaign average",
            2: "Said yes to the AI agent (other ICPs), or asked for a callback",
            3: "Spoke to the AI agent, no yes yet",
            4: "The AI agent never reached them"}.items()],
        [], ["ICP", "Leads", "Lead → sale (last measured)", "Open with", "Story", "Day"],
        *[[k, v["leads"], round(v["sales"] / v["leads"], 4) if v["leads"] else 0,
           icp_map.get("playbook", {}).get(k, {}).get("open_with", ""),
           icp_map.get("playbook", {}).get(k, {}).get("story", ""),
           icp_map.get("playbook", {}).get(k, {}).get("day", "")]
          for k, v in sorted(result["icp"].items(), key=lambda x: -x[1]["leads"]) if v["leads"]],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 36
    for col, w in zip("BCDEF", [12, 14, 70, 60, 8]):
        ws.column_dimensions[col].width = w
    widths = [16, 34, 20, 13, 26, 26, 24, 60, 50, 10, 30, 70, 22, 9, 18, 26, 26, 26, 40, 16, 12, 24, 24]
    for tab in ("Call list", "Already enrolled", "Do not call"):
        t = wb.create_sheet(tab)
        cols = COLS + (["Enrolled status"] if tab == "Already enrolled" else [])
        t.append(cols)
        for r in rows:
            if r["tab"] == tab:
                t.append([r[c] for c in cols])
        for c in t[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        t.freeze_panes = "D2"
        t.auto_filter.ref = t.dimensions
        for i, w in enumerate(widths + [16], start=1):
            t.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in t.iter_rows(min_row=2):
            for c in row[7:12]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
