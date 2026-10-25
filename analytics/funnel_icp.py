"""Leads, TimePay results and sales per ICP for a bootcamp campaign, read from the GrowthX funnels.

Each person (last 10 phone digits, else email) is credited to the first funnel they filled, and that
funnel's landing page is mapped to one ICP in ``analytics/funnel_icp_maps/<bootcamp>.json``. The ICPs
there are the ones the TimePay discovery agent uses, plus "Teachers / educators" and "General"; General
pages also carry a segment (who the page targets) so General can be split without new ICPs.

TimePay: each person's best call outcome in the given campaign(s). Positive = the registration
dispositions in ``POSITIVE`` (Full Conversion, Bootcamp Attendance Confirmed, Bootcamp Confirmed –
Discovery Incomplete, Discovery Successful).

Sales: rows in the enrolment sheet's bootcamp tab whose status is in ``SALE_STATUSES``, matched to a
GrowthX lead by phone, then email.

Every run writes a Checks tab: leads, sales and money must reconcile three ways, every page must be
mapped, and TimePay must return every log it counts. Any FAIL stops the report (exit 3).

New campaigns bring new landing pages: ``pages`` lists each unmapped page with its text so it can be read
and added to the map before ``report`` runs.

    python -m analytics.funnel_icp pages  <bootcamp> <leadtype> <out_dir> [--from YYYY-MM-DD --to YYYY-MM-DD]
    python -m analytics.funnel_icp report <bootcamp> <leadtype> <sales.xlsx> <tab> <out_dir>
        [--from .. --to ..] [--timepay CAMPAIGN_ID[,ID]] [--tp-from YYYY-MM-DDTHH:MM:SS --tp-to ..]
        [--timepay-logs logs.json] [--stated-icp phone_code.tsv]

Stage 1 (managers): ``ICP_funnel_report_<leadtype>.xlsx``. Stage 2 (sales callers): ``Caller_list_<leadtype>.xlsx``,
one row per lead (see ``analytics/funnel_icp_callers.py``).
"""

from __future__ import annotations

import argparse
import html
import json
import os
import random
import re
import sys
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from typing import Any, Iterable

import requests

MAPS_DIR = os.path.join(os.path.dirname(__file__), "funnel_icp_maps")
SALE_STATUSES = {"Disbursed", "One shot", "Installments", "Booking fees received"}
UNMAPPED = "UNMAPPED"
GENERAL = "General"
# Best first. A person's outcome is the best disposition across all their calls.
POSITIVE = ("Full Conversion", "Bootcamp Attendance Confirmed", "Bootcamp Confirmed", "Discovery Successful")


def phone10(value: Any) -> str:
    if isinstance(value, float):
        value = int(value)
    return re.sub(r"\D", "", str(value or ""))[-10:]


def email_key(value: Any) -> str:
    return str(value or "").strip().lower()


def page_slug(url: str | None) -> str:
    path = (url or "").split("?")[0].split("#")[0]
    return path.split("/f/", 1)[1].strip("/") if "/f/" in path else ""


def load_map(bootcamp: str) -> dict:
    with open(os.path.join(MAPS_DIR, f"{bootcamp}.json"), encoding="utf-8") as f:
        return json.load(f)


def person_key(row: dict) -> str:
    return phone10(row.get("whatsapp") or row.get("phone")) or email_key(row.get("email"))


def first_touch(rows: Iterable[dict]) -> dict[str, dict]:
    """Each person's earliest lead row by ``captured_ist``."""
    best: dict[str, dict] = {}
    for row in rows:
        key = person_key(row)
        if key and (key not in best or (row.get("captured_ist") or "") < (best[key].get("captured_ist") or "")):
            best[key] = row
    return best


def outcome_rank(disposition: str | None) -> int:
    """0 = Full Conversion … len(POSITIVE)-1; len(POSITIVE) = connected but not positive."""
    d = disposition or ""
    return next((i for i, p in enumerate(POSITIVE) if d.startswith(p)), len(POSITIVE))


def _text(v: Any) -> str:
    return "" if v in (None, "None", "null") else str(v).strip()


def timepay_by_phone(logs: Iterable[dict]) -> dict[str, dict]:
    """Per phone, each call_id once: dialled, connected (any call Contacted), best outcome rank, attempts,
    last call time (IST), the best call's disposition and the most informative summary of a connected call."""
    seen: set = set()
    out: dict[str, dict] = {}
    for r in logs:
        cid = r.get("call_id")
        if cid in seen:
            continue
        seen.add(cid)
        p = phone10(r.get("to"))
        if not p:
            continue
        s = out.setdefault(p, {"connected": False, "best": len(POSITIVE) + 1, "attempts": 0, "last_call": "",
                               "disposition": "", "summary": "", "dispositions": [], "_score": (99, 0)})
        s["attempts"] += 1
        s["last_call"] = max(s["last_call"], _text(r.get("start_time")))
        if r.get("status") == "Contacted":
            s["connected"] = True
            rank = outcome_rank(r.get("disposition"))
            s["dispositions"].append(_text(r.get("disposition")).replace(" NEW", ""))
            s["best"] = min(s["best"], rank)
            summary = _text(r.get("summary"))
            score = (rank, -len(summary))
            if score < s["_score"]:
                s["_score"] = score
                s["disposition"] = _text(r.get("disposition")).replace(" NEW", "")
                s["summary"] = summary
    for s in out.values():
        s.pop("_score", None)
    return out


def read_sales(xlsx_path: str, tab: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = wb[tab].iter_rows(values_only=True)
    head = [str(h or "").strip().lower() for h in next(rows)]

    def col(prefix: str) -> int:
        return next(i for i, h in enumerate(head) if h.startswith(prefix))

    c_name, c_phone, c_email = col("student name"), col("phone"), col("email")
    c_paid, c_status = col("money received"), col("status of the lead")
    out = []
    for n, r in enumerate(rows, start=2):
        if not r or not r[c_name] or str(r[c_status] or "").strip() not in SALE_STATUSES:
            continue
        paid = r[c_paid]
        out.append({"row": n, "name": str(r[c_name]).strip(), "status": str(r[c_status]).strip(),
                    "phone": phone10(r[c_phone]), "email": email_key(r[c_email]),
                    "paid": float(paid) if isinstance(paid, (int, float)) else 0.0,
                    "paid_raw": paid})
    return out


def _blank() -> dict:
    return {"leads": 0, "dialled": 0, "connected": 0, "positive": 0, "full": 0, "sales": 0, "paid": 0.0}


def build(leads: list[dict], sales: list[dict], icp_map: dict, tp: dict[str, dict] | None = None) -> dict:
    """Aggregate first-touch leads, TimePay outcomes and matched sales per ICP, segment and funnel."""
    pages, segs = icp_map["pages"], icp_map.get("general_segments", {})
    people = first_touch(leads)
    by_email: dict[str, str] = {}
    for k, r in people.items():
        e = email_key(r.get("email"))
        if e:
            by_email.setdefault(e, k)

    def labels(row: dict) -> tuple[str, str, tuple]:
        slug = page_slug(row.get("pageUrl"))
        icp = pages.get(slug, UNMAPPED)
        seg = segs.get(slug, "") if icp == GENERAL else ""
        return icp, seg, (row.get("funnel") or "", slug, icp, seg)

    icp = {name: _blank() for name in icp_map["icps"]}
    seg_agg: dict[str, dict] = defaultdict(_blank)
    funnels: dict[tuple, dict] = defaultdict(_blank)
    unmapped: Counter = Counter()

    def buckets(row: dict) -> list[dict]:
        name, seg, fkey = labels(row)
        out = [icp.setdefault(name, _blank()), funnels[fkey]]
        if name == GENERAL:
            out.append(seg_agg[seg or "(no segment)"])
        return out

    for key, row in people.items():
        if labels(row)[0] == UNMAPPED:
            unmapped[(row.get("pageUrl") or "(no page)").split("?")[0]] += 1
        t = (tp or {}).get(key)
        for b in buckets(row):
            b["leads"] += 1
            if t:
                b["dialled"] += 1
                b["connected"] += t["connected"]
                b["positive"] += t["best"] < len(POSITIVE)
                b["full"] += t["best"] == 0
    matched, sale_rows = Counter(), []
    for sale in sales:
        if sale["phone"] and sale["phone"] in people:
            key, how = sale["phone"], "phone"
        elif sale["email"] and sale["email"] in by_email:
            key, how = by_email[sale["email"]], "email"
        else:
            key, how = None, "not in GrowthX"
        matched[how] += 1
        row = people.get(key) if key else None
        sale_rows.append({**sale, "how": how, "key": key,
                          "funnel": row.get("funnel") if row else "", "slug": page_slug(row.get("pageUrl")) if row else "",
                          "icp": labels(row)[0] if row else "", "segment": labels(row)[1] if row else "",
                          "captured_ist": row.get("captured_ist") if row else ""})
        if row:
            for b in buckets(row):
                b["sales"] += 1
                b["paid"] += sale["paid"]
    return {"icp": icp, "segments": dict(seg_agg), "funnels": dict(funnels), "unmapped": dict(unmapped),
            "matched": dict(matched), "people": len(people), "raw_leads": len(leads), "sales": sale_rows,
            "sales_total": len(sales), "timepay": tp is not None,
            "tp_numbers": len(tp or {}), "tp_in_growthx": len(set(tp or {}) & set(people))}


def checks(result: dict, icp_map: dict, leads: list[dict], tp_expected: int | None = None,
           tp_read: int | None = None) -> list[tuple[str, str, str]]:
    """(status, check, detail). FAIL stops the report; WARN needs a look; INFO is context."""
    out: list[tuple[str, str, str]] = []

    def add(ok: bool, name: str, detail: str, level: str = "FAIL") -> None:
        out.append(("PASS" if ok else level, name, detail))

    icps, segs = set(icp_map["icps"]), icp_map.get("general_segments", {})
    bad = {s: v for s, v in icp_map["pages"].items() if v not in icps}
    add(not bad, "Every mapped page uses an ICP from the agent list", str(bad) if bad else "ok")
    gen = {s for s, v in icp_map["pages"].items() if v == GENERAL}
    add(gen == set(segs), "Every General page has a segment, and only General pages",
        f"missing {sorted(gen - set(segs))}, extra {sorted(set(segs) - gen)}")
    add(not result["unmapped"], "Every lead's landing page is mapped",
        f"{sum(result['unmapped'].values())} leads on {len(result['unmapped'])} pages" if result["unmapped"] else "ok")
    for field in ("leads", "dialled", "connected", "positive", "full", "sales", "paid"):
        a = sum(v[field] for v in result["icp"].values())
        b = sum(v[field] for v in result["funnels"].values())
        add(round(a, 2) == round(b, 2), f"{field}: ICP total = funnel total", f"{a:,.0f} vs {b:,.0f}")
    ga = result["icp"].get(GENERAL, _blank())
    for field in ("leads", "sales", "paid"):
        s = sum(v[field] for v in result["segments"].values())
        add(round(s, 2) == round(ga[field], 2), f"General {field} = sum of its segments", f"{s:,.0f} vs {ga[field]:,.0f}")
    add(sum(v["leads"] for v in result["icp"].values()) == result["people"], "Leads = unique people in GrowthX",
        f"{result['people']:,} people from {result['raw_leads']:,} form fills")
    add(sum(result["matched"].values()) == result["sales_total"], "Every sale row counted once (matched + unmatched)",
        f"{result['matched']} of {result['sales_total']}")
    paid_matched = sum(s["paid"] for s in result["sales"] if s["key"])
    add(round(paid_matched, 2) == round(sum(v["paid"] for v in result["icp"].values()), 2),
        "Money received of matched sales = ICP total", f"Rs {paid_matched:,.0f}")
    nonnum = [s["row"] for s in result["sales"] if s["paid_raw"] not in (None, "") and not isinstance(s["paid_raw"], (int, float))]
    add(not nonnum, "Money received is a number on every sale row", f"sheet rows {nonnum[:20]}", "WARN")
    ids = Counter(s["phone"] or s["email"] for s in result["sales"] if s["phone"] or s["email"])
    dups = sorted(k[-4:] for k, n in ids.items() if n > 1)
    add(not dups, "No person appears twice as a sale", f"{len(dups)} repeated (phone ends {dups[:10]})", "WARN")
    blank = [s["row"] for s in result["sales"] if not s["phone"] and not s["email"]]
    add(not blank, "Every sale row has a phone or email", f"sheet rows {blank}", "WARN")
    em = [s["row"] for s in result["sales"] if s["how"] == "email" and s["phone"]]
    add(not em, "Email-only matches where the sheet has a different phone", f"sheet rows {em} (see Spot check)", "WARN")
    no_key = sum(1 for r in leads if not person_key(r))
    add(no_key == 0, "Every GrowthX lead has a phone or email", f"{no_key} leads skipped", "WARN")
    if tp_expected is not None:
        add(tp_read == tp_expected, "TimePay: every counted call was read", f"read {tp_read:,} of {tp_expected:,}")
    if result["timepay"]:
        out.append(("INFO", "TimePay numbers found in this GrowthX leadtype",
                    f"{result['tp_in_growthx']:,} of {result['tp_numbers']:,} dialled numbers"))
    return out


def page_text(url: str, timeout: float = 20) -> str:
    try:
        body = requests.get(url, timeout=timeout).text
    except requests.RequestException as exc:
        return f"(could not open: {exc.__class__.__name__})"
    body = re.sub(r"<script.*?</script>|<style.*?</style>", " ", body, flags=re.S | re.I)
    return " ".join(html.unescape(re.sub(r"<[^>]+>", " ", body)).split())[:700]


def fetch_leads(leadtype: str, start: str | None, end: str | None) -> list[dict]:
    import integrations.env  # noqa: F401  loads .env
    from integrations.growthx.client import GrowthXClient, captured_ist

    client = GrowthXClient()
    if start and end:
        return client.leads_for_ist_days(start, end, leadtype=leadtype)[0]
    rows = []
    for row in client.leads(leadtype=leadtype):
        at = captured_ist(row.get("capturedAt"))
        rows.append({**row, "captured_ist": at.strftime("%Y-%m-%dT%H:%M") if at else ""})
    return rows


def fetch_timepay(campaign_ids: list[str], start: str, end: str) -> tuple[list[dict], int]:
    """All call logs of the campaigns in a fixed window (so new calls cannot shift pages), and the count
    TimePay reports for it. Read is deduped by call_id."""
    import integrations.env  # noqa: F401
    from integrations.timepay import TimePayClient

    tp = TimePayClient()
    rows, expected = [], 0
    for cid in campaign_ids:
        params = {"type": "call", "campaign_id": cid, "start_time": start, "end_time": end}
        n = tp.count_logs(start, end, type="call", campaign_id=cid)
        expected += n

        def page(p: int) -> list[dict]:
            err: Exception | None = None
            for _ in range(4):
                try:
                    return tp.get("logs", {**params, "page": p}).get("data", [])
                except Exception as exc:  # retried; the last error is raised
                    err = exc
            raise err  # type: ignore[misc]

        with ThreadPoolExecutor(8) as ex:
            for chunk in ex.map(page, range(1, -(-n // 10) + 1)):
                rows.extend(chunk)
    return rows, expected


def _pct(a: float, b: float) -> float:
    return round(a / b, 4) if b else 0.0


def write_xlsx(result: dict, chk: list, path: str, title: str, fresh: str) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "By ICP"
    ws.append([title])
    ws.append([fresh])
    ws.append([f"{result['people']:,} unique leads, each credited to their first funnel. {result['sales_total']} "
               f"sales in the sheet; matched {result['matched']}. Positive = FC, BAC, BC-DI or Discovery Successful."])
    ws.append([])
    head = ["ICP", "General segment", "Leads", "Share of leads"]
    if result["timepay"]:
        head += ["TimePay dialled", "Connected", "Connect %", "Positive", "Positive % of connected",
                 "Full Conversion"]
    head += ["Sales", "Lead → sale %", "Money received (Rs)"]
    ws.append(head)

    def line(icp: str, seg: str, a: dict) -> list:
        r = [icp, seg, a["leads"], _pct(a["leads"], result["people"])]
        if result["timepay"]:
            r += [a["dialled"], a["connected"], _pct(a["connected"], a["dialled"]), a["positive"],
                  _pct(a["positive"], a["connected"]), a["full"]]
        return r + [a["sales"], _pct(a["sales"], a["leads"]), round(a["paid"])]

    for name, a in sorted(result["icp"].items(), key=lambda x: -x[1]["leads"]):
        if not a["leads"] and not a["sales"]:
            continue
        ws.append(line(name, "All" if name == GENERAL else "", a))
        if name == GENERAL:
            for seg, s in sorted(result["segments"].items(), key=lambda x: -x[1]["leads"]):
                ws.append(line("", seg, s))
    fs = wb.create_sheet("By funnel")
    fs.append(["Funnel", "Landing page", "ICP", "General segment"] + head[2:])
    for (funnel, slug, name, seg), a in sorted(result["funnels"].items(), key=lambda x: -x[1]["leads"]):
        fs.append([funnel, slug, name, seg] + line(name, seg, a)[2:])
    cs = wb.create_sheet("Checks")
    cs.append(["Result", "Check", "Detail"])
    for c in chk:
        cs.append(list(c))
    sp = wb.create_sheet("Spot check")
    sp.append(["Check each by hand in the sheet and GrowthX before sharing.", "", "", "", "", "", "", ""])
    sp.append(["Sheet row", "Student", "Status", "Matched by", "First funnel", "Landing page", "ICP", "Segment"])
    picks = [s for s in result["sales"] if s["how"] == "email" and s["phone"]]
    rest = [s for s in result["sales"] if s not in picks]
    random.Random(7).shuffle(rest)
    for s in (picks + rest)[:max(20, len(picks))]:
        sp.append([s["row"], s["name"], s["status"], s["how"], s["funnel"], s["slug"], s["icp"], s["segment"]])
    if result["unmapped"]:
        us = wb.create_sheet("Unmapped pages")
        us.append(["Landing page", "Leads"])
        for url, n in sorted(result["unmapped"].items(), key=lambda x: -x[1]):
            us.append([url, n])
    wb.save(path)


def verify_written(path: str, result: dict) -> bool:
    """Re-open the saved file and confirm the ICP rows add up to the same leads and sales."""
    import openpyxl

    ws = openpyxl.load_workbook(path, read_only=True)["By ICP"]
    rows = list(ws.iter_rows(min_row=6, values_only=True))
    head = list(next(openpyxl.load_workbook(path, read_only=True)["By ICP"].iter_rows(min_row=5, max_row=5,
                                                                                         values_only=True)))
    top = [r for r in rows if r[0]]
    leads = sum(r[head.index("Leads")] for r in top)
    sales = sum(r[head.index("Sales")] for r in top)
    return leads == result["people"] and sales == sum(v["sales"] for v in result["icp"].values())


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("pages", "report"):
        p = sub.add_parser(name)
        p.add_argument("bootcamp")
        p.add_argument("leadtype")
        if name == "report":
            p.add_argument("sales_xlsx")
            p.add_argument("tab")
        p.add_argument("out_dir")
        p.add_argument("--from", dest="start")
        p.add_argument("--to", dest="end")
        if name == "report":
            p.add_argument("--timepay", help="TimePay campaign id(s), comma separated")
            p.add_argument("--tp-from", default="2026-01-01T00:00:00")
            p.add_argument("--tp-to")
            p.add_argument("--timepay-logs", help="saved logs JSON instead of fetching")
            p.add_argument("--stated-icp", help="TSV of phone<TAB>code: ICP read from each lead's call summary")
    a = ap.parse_args(argv)
    icp_map = load_map(a.bootcamp)
    os.makedirs(a.out_dir, exist_ok=True)
    leads = fetch_leads(a.leadtype, a.start, a.end)

    if a.cmd == "pages":
        seen: dict[str, dict] = {}
        for row in leads:
            slug = page_slug(row.get("pageUrl"))
            if slug and slug not in icp_map["pages"]:
                e = seen.setdefault(slug, {"url": row["pageUrl"].split("?")[0], "funnels": Counter(), "leads": 0})
                e["leads"] += 1
                e["funnels"][row.get("funnel") or ""] += 1
        out = [{"slug": s, "leads": e["leads"], "funnels": dict(e["funnels"]), "url": e["url"],
                "page_text": page_text(e["url"])} for s, e in sorted(seen.items(), key=lambda x: -x[1]["leads"])]
        no_page = sum(1 for r in leads if not page_slug(r.get("pageUrl")))
        path = os.path.join(a.out_dir, f"unmapped_pages_{a.leadtype}.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(out, f, indent=1, ensure_ascii=False)
        print(f"{len(leads)} leads, {len(out)} landing pages not in {a.bootcamp}.json, "
              f"{no_page} leads with no landing page -> {path}")
        return 0

    from integrations.timeutil import IST
    from datetime import datetime

    tp, tp_expected, tp_read = None, None, None
    if a.timepay_logs:
        with open(a.timepay_logs, encoding="utf-8") as f:
            logs = json.load(f)
        tp = timepay_by_phone(logs)
    elif a.timepay:
        end = a.tp_to or datetime.now(IST).strftime("%Y-%m-%dT%H:%M:%S")
        logs, tp_expected = fetch_timepay(a.timepay.split(","), a.tp_from, end)
        tp_read = len({r.get("call_id") for r in logs})
        tp = timepay_by_phone(logs)
    result = build(leads, read_sales(a.sales_xlsx, a.tab), icp_map, tp)
    chk = checks(result, icp_map, leads, tp_expected, tp_read)
    from analytics.funnel_icp_callers import caller_checks, caller_rows, read_stated, write_caller_xlsx

    callers = caller_rows(leads, result, icp_map, tp, read_stated(a.stated_icp, icp_map.get("stated_icp_codes", {})))
    chk += caller_checks(callers, result)
    path = os.path.join(a.out_dir, f"ICP_funnel_report_{a.leadtype}.xlsx")
    fresh = f"Data pulled {datetime.now(IST):%d %b %Y, %I:%M %p} IST"
    write_xlsx(result, chk, path, f"Leads, TimePay and sales per ICP · {a.bootcamp} · {a.leadtype}", fresh)
    ok_file = verify_written(path, result)
    chk.append(("PASS" if ok_file else "FAIL", "Saved file re-read: ICP rows add up", path))
    write_xlsx(result, chk, path, f"Leads, TimePay and sales per ICP · {a.bootcamp} · {a.leadtype}", fresh)
    caller_path = os.path.join(a.out_dir, f"Caller_list_{a.leadtype}.xlsx")
    write_caller_xlsx(callers, caller_path, f"Caller list · {a.bootcamp} · {a.leadtype}", fresh, icp_map, result)
    for c in chk:
        if c[0] != "PASS":
            print(f"{c[0]:5} {c[1]}: {c[2]}")
    fails = [c for c in chk if c[0] == "FAIL"]
    print(f"{len(chk) - len(fails)} of {len(chk)} checks not failed -> {path} and {caller_path}")
    if fails:
        print("REPORT NOT READY: fix the FAIL checks above before sharing.", file=sys.stderr)
        return 3
    for name, v in sorted(result["icp"].items(), key=lambda x: -x[1]["leads"]):
        if v["leads"] or v["sales"]:
            tpx = (f" dialled {v['dialled']:6} conn {v['connected']:6} pos {v['positive']:5} "
                   f"({100 * _pct(v['positive'], v['connected']):4.1f}%)") if tp is not None else ""
            print(f"{name:40} {v['leads']:6} leads{tpx} {v['sales']:4} sales "
                  f"{100 * _pct(v['sales'], v['leads']):4.1f}%  Rs {v['paid']:,.0f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
